"""Contracts for embedding fonts into a .docx.

A font name is not a font. The document carries a family name and the
reader's machine resolves it, which for Arabic means tofu on any machine
without the face installed. docx_embed_fonts.py puts the font bytes in
the package, in parts python-docx does not model, so the things worth
testing above everything else are a package that is still internally
consistent, a document that still opens, and bytes that come back out of
the obfuscation exactly as they went in.

The obfuscation is the part with a trap in it. XOR is its own inverse, so
a round trip through the script's own key derivation passes whether or
not the key byte order matches the specification, and a file Word would
refuse would still look green here. So the order is pinned twice over: a
known answer vector written out from the specification independently of
the script, and, where LibreOffice can produce one, a .odttf written by
another implementation entirely.

The third thing is the licence gate. Embedding ships the font inside
every copy of the deliverable, so a face with no licence file on disk is
refused unless the caller says otherwise, and that refusal has to keep
working.

Most tests build their own font directory out of one .ttf copied under a
synthetic family name, with and without an OFL.txt beside it. That keeps
the gate deterministic instead of depending on which faces this
particular machine happens to have and where its distribution files
their licences.
"""

from __future__ import annotations

import json
import os
import shutil
import subprocess
import sys
import zipfile
from pathlib import Path

import pytest

REPO = Path(__file__).resolve().parents[2]
CLI = REPO / "skills" / "productivity" / "docx" / "scripts" / \
    "docx_embed_fonts.py"
CREATE = REPO / "skills" / "productivity" / "docx" / "scripts" / \
    "docx_create.py"

W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
CT_NS = "http://schemas.openxmlformats.org/package/2006/content-types"
REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
CT_OBFUSCATED = ("application/vnd.openxmlformats-officedocument"
                 ".obfuscatedFont")

SLOTS = {"regular": "embedRegular", "bold": "embedBold",
         "italic": "embedItalic", "boldItalic": "embedBoldItalic"}

FAMILY = "HermesTestFace"
ARABIC_FAMILY = "IBM Plex Sans Arabic"
ARABIC_TEXT = "الإيرادات ارتفعت بنسبة 12 بالمئة"


def run(*args: str, fonts: Path | None = None) -> dict:
    out = _invoke(args, fonts)
    assert out.returncode == 0, f"{args}: {out.stderr}"
    return json.loads(out.stdout)


def _invoke(args, fonts):
    env = dict(os.environ)
    # Both variables are popped, not just the one under test: the shared
    # search path reads PPTX_FONT_DIRS, so a stray one in the ambient
    # environment would quietly widen the search this test is pinning.
    env.pop("PPTX_FONT_DIRS", None)
    env.pop("DOCX_FONT_DIRS", None)
    if fonts is not None:
        env["DOCX_FONT_DIRS"] = str(fonts)
    return subprocess.run([sys.executable, str(CLI), *args],
                          capture_output=True, text=True, env=env)


# ---------------------------------------------------------------------------
# Fixtures
# ---------------------------------------------------------------------------


def _some_ttf() -> Path | None:
    """Any TrueType face on this machine, to copy under a test name.

    The bytes only have to be a real font; nothing in the embedding path
    reads the family out of the file, so a copy renamed to the synthetic
    family exercises exactly the code a real face would.
    """
    try:
        out = subprocess.run(["fc-match", "-f", "%{file}", "sans"],
                             capture_output=True, text=True, timeout=10)
    except (OSError, subprocess.SubprocessError):
        return None
    path = Path(out.stdout.strip()) if out.returncode == 0 else None
    if path is not None and path.suffix.lower() == ".ttf" and path.is_file():
        return path
    for base in (Path("/usr/share/fonts"), Path.home() / ".local/share/fonts"):
        if base.is_dir():
            for hit in base.rglob("*.ttf"):
                return hit
    return None


@pytest.fixture
def ttf():
    path = _some_ttf()
    if path is None:
        pytest.skip("no .ttf on this machine to build a test family from")
    return path


@pytest.fixture
def licensed_fonts(tmp_path, ttf):
    """A font directory holding the test family and its OFL licence."""
    root = tmp_path / "licensed"
    (root / FAMILY).mkdir(parents=True)
    shutil.copy(ttf, root / FAMILY / f"{FAMILY}-Regular.ttf")
    shutil.copy(ttf, root / FAMILY / f"{FAMILY}-Bold.ttf")
    (root / FAMILY / "OFL.txt").write_text(
        "SIL OPEN FONT LICENSE Version 1.1\n", encoding="utf-8")
    return root


@pytest.fixture
def unlicensed_fonts(tmp_path, ttf):
    """The same family with no licence file anywhere above it."""
    root = tmp_path / "unlicensed"
    (root / FAMILY).mkdir(parents=True)
    shutil.copy(ttf, root / FAMILY / f"{FAMILY}-Regular.ttf")
    return root


def _build(tmp_path, name, spec) -> Path:
    """Build a document with the skill's own creator, from a JSON spec."""
    pytest.importorskip("docx")
    spec_path = tmp_path / (name + ".json")
    spec_path.write_text(json.dumps(spec, ensure_ascii=False),
                         encoding="utf-8")
    out = tmp_path / name
    result = subprocess.run(
        [sys.executable, str(CREATE), str(spec_path), str(out)],
        capture_output=True, text=True)
    assert result.returncode == 0, result.stderr
    return out


def _with_family(tmp_path, name, family, text) -> Path:
    """A document whose body style names one family outright.

    docx_create.py takes its fonts from the house theme, so the family is
    set afterwards through python-docx: the point of the fixture is a
    document that asks for a known family, by whatever route.
    """
    from docx import Document

    path = _build(tmp_path, name, {"blocks": [
        {"type": "heading", "text": text, "level": 1},
        {"type": "paragraph", "text": text},
    ]})
    if family:
        document = Document(str(path))
        for paragraph in document.paragraphs:
            for run_ in paragraph.runs:
                run_.font.name = family
                # w:cs is the one Arabic text actually reads, and
                # python-docx sets only w:ascii and w:hAnsi.
                rfonts = run_._element.get_or_add_rPr().get_or_add_rFonts()
                rfonts.set(f"{{{W_NS}}}cs", family)
        document.save(str(path))
    return path


@pytest.fixture
def doc(tmp_path):
    """A document whose runs name the synthetic test family."""
    return _with_family(tmp_path, "doc.docx", FAMILY, "Quarter three")


@pytest.fixture
def latin_doc(tmp_path):
    """A document that names no family of its own, only the house's."""
    return _with_family(tmp_path, "latin.docx", None,
                        "Handles carried the quarter")


# ---------------------------------------------------------------------------
# Helpers that read the package the way Word would
# ---------------------------------------------------------------------------


def package(path: Path) -> dict:
    from lxml import etree

    with zipfile.ZipFile(path) as z:
        names = set(z.namelist())
        blobs = {n: z.read(n) for n in names}
    types = etree.fromstring(blobs["[Content_Types].xml"])
    overrides = {ov.get("PartName"): ov.get("ContentType")
                 for ov in types.findall(f"{{{CT_NS}}}Override")}
    defaults = {(d.get("Extension") or "").lower(): d.get("ContentType")
                for d in types.findall(f"{{{CT_NS}}}Default")}
    table = etree.fromstring(blobs["word/fontTable.xml"]) \
        if "word/fontTable.xml" in names else None
    settings = etree.fromstring(blobs["word/settings.xml"]) \
        if "word/settings.xml" in names else None
    targets = {}
    # The r:id in w:embedRegular resolves against the font table's own
    # rels, which is the structural detail most easily got wrong.
    rels_name = "word/_rels/fontTable.xml.rels"
    if rels_name in names:
        rels = etree.fromstring(blobs[rels_name])
        for rel in rels.findall(f"{{{REL_NS}}}Relationship"):
            targets[rel.get("Id")] = "word/" + rel.get("Target").lstrip("/")
    return {"names": names, "blobs": blobs, "table": table,
            "settings": settings, "overrides": overrides,
            "defaults": defaults, "targets": targets}


def content_type_of(pkg, part):
    override = pkg["overrides"].get("/" + part)
    if override:
        return override
    return pkg["defaults"].get(part.rsplit(".", 1)[-1].lower())


def embedded_entries(pkg) -> dict:
    """Family to slot to (relationship id, font key), out of the XML."""
    out = {}
    if pkg["table"] is None:
        return out
    for font in pkg["table"].findall(f"{{{W_NS}}}font"):
        slots = {}
        for slot, tag in SLOTS.items():
            element = font.find(f"{{{W_NS}}}{tag}")
            if element is not None:
                slots[slot] = (element.get(f"{{{R_NS}}}id"),
                               element.get(f"{{{W_NS}}}fontKey"))
        if slots:
            out[font.get(f"{{{W_NS}}}name")] = slots
    return out


def deobfuscate(data: bytes, font_key: str) -> bytes:
    """The specification's rule, written out here and not imported.

    Strip the braces and hyphens from the GUID, read the 32 hex digits as
    16 bytes, reverse them, and XOR them over the first 32 bytes of the
    part. Spelling it a second time is the point: importing the script's
    own function would make the round trip pass for any byte order, XOR
    being its own inverse.
    """
    digits = font_key.strip().strip("{}").replace("-", "")
    assert len(digits) == 32, font_key
    key = bytes(reversed(bytes.fromhex(digits)))
    out = bytearray(data)
    for i in range(32):
        out[i] ^= key[i % 16]
    return bytes(out)


def assert_consistent(path: Path):
    """Every reference resolves, every part is typed, nothing orphaned."""
    pkg = package(path)
    referenced = set()
    for family, slots in embedded_entries(pkg).items():
        for slot, (rel_id, font_key) in slots.items():
            target = pkg["targets"].get(rel_id)
            assert target, \
                f"{family} {slot}: {rel_id} is not in fontTable.xml.rels"
            assert target in pkg["names"], f"{target} is not in the package"
            assert content_type_of(pkg, target) == CT_OBFUSCATED
            assert font_key and font_key.startswith("{")
            plain = deobfuscate(pkg["blobs"][target], font_key)
            assert plain[:4] in (b"\x00\x01\x00\x00", b"OTTO", b"true",
                                 b"ttcf"), \
                f"{family} {slot}: {target} is not a font once de-obfuscated"
            referenced.add(target)
    for name in pkg["names"]:
        if name.startswith("word/fonts/") and name.endswith(".odttf"):
            assert name in referenced, f"{name} is orphaned"

    if referenced:
        # CT_Font fixes the order of its children and the four embed
        # elements go last. Out of place is what makes Word offer to
        # repair the file.
        order = ("altName", "panose1", "charset", "family", "notTrueType",
                 "pitch", "sig", "embedRegular", "embedBold", "embedItalic",
                 "embedBoldItalic")
        for font in pkg["table"].findall(f"{{{W_NS}}}font"):
            seen = [child.tag.split("}")[-1] for child in font]
            ranked = [order.index(name) for name in seen if name in order]
            assert ranked == sorted(ranked), seen


def reopens(path: Path):
    from docx import Document

    return [p.text for p in Document(str(path)).paragraphs if p.text]


# ---------------------------------------------------------------------------
# The obfuscation
# ---------------------------------------------------------------------------

# A known answer vector, so the key byte order is pinned by something
# other than the script agreeing with itself. The key below is the GUID
# read as 16 bytes and reversed, and the expected output is that key
# XORed over 32 zero bytes, which is just the key twice.
VECTOR_KEY = "{9E7C1B3A-4D5F-6071-8293-A4B5C6D7E8F9}"
VECTOR_PLAIN = bytes(32)
VECTOR_CIPHER = bytes.fromhex(
    "f9e8d7c6b5a493827160" "5f4d" "3a1b7c9e"
    "f9e8d7c6b5a493827160" "5f4d" "3a1b7c9e")


def test_the_key_byte_order_is_the_documented_one():
    """Reversed GUID bytes, XORed over the first 32 bytes only."""
    sys.path.insert(0, str(CLI.parent))
    try:
        import docx_embed_fonts
    finally:
        sys.path.pop(0)
    assert docx_embed_fonts.obfuscate(VECTOR_PLAIN, VECTOR_KEY) \
        == VECTOR_CIPHER
    assert deobfuscate(VECTOR_CIPHER, VECTOR_KEY) == VECTOR_PLAIN
    # Byte 32 onwards is untouched, which is why a whole font file can be
    # recovered from the part.
    longer = VECTOR_PLAIN + b"\xAB" * 16
    assert docx_embed_fonts.obfuscate(longer, VECTOR_KEY)[32:] == b"\xAB" * 16


def test_the_stored_part_is_the_source_font_byte_for_byte(
        doc, licensed_fonts):
    """The whole point: what Word reads back has to be the font.

    De-obfuscated with the key out of the XML and compared against the
    file on disk, not against anything the script kept in memory.
    """
    report = run("embed", str(doc), fonts=licensed_fonts)
    record = [r for r in report["embedded"] if r["family"] == FAMILY][0]
    pkg = package(doc)
    entries = embedded_entries(pkg)[FAMILY]
    assert set(record["slots"]) == set(entries)
    for slot, info in record["slots"].items():
        rel_id, font_key = entries[slot]
        stored = pkg["blobs"][pkg["targets"][rel_id]]
        source = Path(info["file"]).read_bytes()
        assert stored != source, "the part is stored obfuscated, not plain"
        assert deobfuscate(stored, font_key) == source


@pytest.mark.skipif(shutil.which("soffice") is None,
                    reason="LibreOffice is not installed")
def test_a_part_written_by_another_implementation_de_obfuscates(tmp_path):
    """Ground truth from outside this repository.

    A known answer vector pins the rule as this test reads the
    specification. This one pins it against a file neither the script nor
    the test wrote: LibreOffice embeds fonts on export to .docx, and its
    .odttf parts have to come back out as fonts under the same rule.
    """
    docx = _libreoffice_embedded_docx(tmp_path)
    if docx is None:
        pytest.skip("this LibreOffice did not embed fonts on export")
    pkg = package(docx)
    entries = embedded_entries(pkg)
    parts = [name for name in pkg["names"]
             if name.startswith("word/fonts/") and name.endswith(".odttf")]
    assert parts
    checked = 0
    for slots in entries.values():
        for rel_id, font_key in slots.values():
            target = pkg["targets"].get(rel_id)
            if target is None:
                continue
            plain = deobfuscate(pkg["blobs"][target], font_key)
            assert plain[:4] in (b"\x00\x01\x00\x00", b"OTTO", b"true",
                                 b"ttcf"), \
                "the key byte order does not match another implementation"
            checked += 1
    assert checked, "no embedded slot was found to check"


def _libreoffice_embedded_docx(tmp_path) -> Path | None:
    """Ask LibreOffice for a .docx with fonts embedded, or None.

    EmbedFonts is a document setting, and it is not honoured out of a
    flat .fodt, so the file is round tripped to .odt first and the
    setting flipped in the package before the export to .docx.
    """
    from lxml import etree  # noqa: F401  (kept for a clear import error)

    fodt = tmp_path / "seed.fodt"
    fodt.write_text(_FODT, encoding="utf-8")
    profile = tmp_path / "loprofile"
    if _soffice(["--convert-to", "odt", "--outdir", str(tmp_path),
                 str(fodt)], profile) is None:
        return None
    odt = tmp_path / "seed.odt"
    if not odt.is_file():
        return None

    unpacked = tmp_path / "odt"
    shutil.unpack_archive(str(odt), str(unpacked), "zip")
    settings = unpacked / "settings.xml"
    text = settings.read_text(encoding="utf-8")
    if "EmbedFonts" not in text:
        return None
    settings.write_text(
        text.replace('name="EmbedFonts" config:type="boolean">false',
                     'name="EmbedFonts" config:type="boolean">true'),
        encoding="utf-8")
    embedded = tmp_path / "embedded.odt"
    with zipfile.ZipFile(embedded, "w", zipfile.ZIP_DEFLATED) as z:
        for item in sorted(unpacked.rglob("*")):
            if item.is_file():
                z.write(item, str(item.relative_to(unpacked)))
    if _soffice(["--convert-to", "docx", "--outdir", str(tmp_path),
                 str(embedded)], profile) is None:
        return None
    out = tmp_path / "embedded.docx"
    if not out.is_file():
        return None
    with zipfile.ZipFile(out) as z:
        if not any(n.endswith(".odttf") for n in z.namelist()):
            return None
    return out


_FODT = """<?xml version="1.0" encoding="UTF-8"?>
<office:document
 xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0"
 xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0"
 xmlns:style="urn:oasis:names:tc:opendocument:xmlns:style:1.0"
 xmlns:fo="urn:oasis:names:tc:opendocument:xmlns:xsl-fo-compatible:1.0"
 office:version="1.3"
 office:mimetype="application/vnd.oasis.opendocument.text">
<office:automatic-styles>
<style:style style:name="P1" style:family="paragraph">
<style:text-properties style:font-name="DejaVu Sans"
 fo:font-family="DejaVu Sans"/></style:style>
</office:automatic-styles>
<office:body><office:text>
<text:p text:style-name="P1">Quarter three</text:p>
</office:text></office:body></office:document>
"""


def _soffice(args, profile: Path):
    """Run LibreOffice headless on a profile of its own.

    Several agents share this machine, and a shared LibreOffice profile
    is a collision, so every run gets a private UserInstallation.
    """
    try:
        out = subprocess.run(
            ["soffice", "--headless", "--norestore",
             f"-env:UserInstallation=file://{profile}", *args],
            capture_output=True, text=True, timeout=300)
    except (OSError, subprocess.SubprocessError):
        return None
    return out if out.returncode == 0 else None


# ---------------------------------------------------------------------------
# What the document uses
# ---------------------------------------------------------------------------


def test_report_names_the_family_the_runs_ask_for(doc, licensed_fonts):
    report = run("report", str(doc), fonts=licensed_fonts)
    assert FAMILY in report["families_used"]
    assert "document" in report["families_used"][FAMILY]
    assert report["dry_run"] is True
    # A dry run writes nothing.
    assert report["bytes_added"] > 0
    assert not [n for n in package(doc)["names"] if n.endswith(".odttf")]


def test_the_walk_reaches_the_styles_and_resolves_the_theme(
        latin_doc, licensed_fonts):
    """A document almost never names a family on a run.

    The house style sets it in styles.xml, and docDefaults points at the
    theme through w:rFonts theme attributes. So the families here have to
    be real names and none of them may be a token.
    """
    report = run("report", str(latin_doc), fonts=licensed_fonts)
    families = report["families_used"]
    assert families, "a built document still asks for the house fonts"
    assert not any(name.startswith("+") for name in families)
    assert any("styles" in where for where in families.values())


def test_the_chart_labels_are_walked_too(tmp_path, licensed_fonts):
    """A chart axis is DrawingML, and an Arabic one is easy to miss."""
    doc = _build(tmp_path, "chart.docx", {"blocks": [
        {"type": "paragraph", "text": ARABIC_TEXT},
        {"type": "chart", "chart": "column", "title": ARABIC_TEXT,
         "categories": ["الربع الأول", "الربع الثاني"],
         "series": {"2025": [12, 18]}},
    ]})
    if "word/charts/chart1.xml" not in package(doc)["names"]:
        pytest.skip("this docx_create.py built no chart part")
    report = run("report", str(doc), fonts=licensed_fonts)
    assert any("charts" in where
               for where in report["families_used"].values()), \
        report["families_used"]


def test_a_document_with_no_arabic_embeds_only_what_it_uses(
        latin_doc, licensed_fonts):
    report = run("embed", str(latin_doc), fonts=licensed_fonts)
    assert FAMILY not in report["families_used"]
    assert all(record["family"] != FAMILY for record in report["embedded"])
    for record in report["embedded"]:
        assert record["family"] in report["families_used"]
    assert_consistent(latin_doc)


# ---------------------------------------------------------------------------
# The licence gate
# ---------------------------------------------------------------------------


def test_licence_gate_refuses_a_face_with_no_licence_file(
        doc, unlicensed_fonts):
    # Restricted to the test family: the house fonts the document also
    # asks for are installed and licensed on this machine, and embedding
    # them would say nothing about the gate.
    report = run("embed", str(doc), "--family", FAMILY,
                 fonts=unlicensed_fonts)
    refused = [s for s in report["skipped"] if s["family"] == FAMILY]
    assert refused, "the unlicensed family should have been refused"
    assert refused[0]["reason"] == "no licence file found"
    assert any("--allow-unlicensed" in line for line in refused[0]["detail"])
    assert not [n for n in package(doc)["names"] if n.endswith(".odttf")]


def test_the_override_flag_embeds_the_same_refused_face(
        doc, unlicensed_fonts):
    report = run("embed", str(doc), "--family", FAMILY,
                 "--allow-unlicensed", fonts=unlicensed_fonts)
    record = [r for r in report["embedded"] if r["family"] == FAMILY]
    assert record, "the override should have let the face through"
    assert record[0]["licence"] is None
    assert record[0]["licence_override"] is True
    assert_consistent(doc)


def test_a_licence_beside_the_font_passes_the_gate_and_is_reported(
        doc, licensed_fonts):
    report = run("embed", str(doc), fonts=licensed_fonts)
    record = [r for r in report["embedded"] if r["family"] == FAMILY][0]
    assert record["licence_override"] is False
    assert record["licence"].endswith("OFL.txt")
    assert Path(record["licence"]).is_file()


# ---------------------------------------------------------------------------
# Missing faces
# ---------------------------------------------------------------------------


def test_a_missing_face_is_reported_and_not_fatal(doc, licensed_fonts):
    report = run("embed", str(doc), "--family", "No Such Family At All",
                 "--allow-unlicensed", fonts=licensed_fonts)
    assert report["ok"] is True
    missing = [s for s in report["skipped"]
               if s["family"] == "No Such Family At All"]
    assert missing and missing[0]["reason"] == "no font file found"
    assert reopens(doc)


def test_one_missing_face_does_not_cost_the_others(doc, licensed_fonts):
    report = run("embed", str(doc), "--family", "No Such Family At All",
                 "--family", FAMILY, fonts=licensed_fonts)
    assert [r["family"] for r in report["embedded"]] == [FAMILY]
    assert [s["family"] for s in report["skipped"]] == \
        ["No Such Family At All"]
    assert_consistent(doc)


# ---------------------------------------------------------------------------
# The package after the pass
# ---------------------------------------------------------------------------


def test_the_package_stays_valid_and_the_document_still_opens(
        doc, licensed_fonts):
    text_before = reopens(doc)
    report = run("embed", str(doc), fonts=licensed_fonts)
    assert report["bytes_added"] > 0
    assert report["size_after"] > report["size_before"]
    assert_consistent(doc)
    assert reopens(doc) == text_before
    assert run("verify", str(doc))["problems"] == []


def test_the_relationships_live_beside_the_font_table(doc, licensed_fonts):
    """An r:id resolves against the rels of the part it appears in.

    w:embedRegular sits in fontTable.xml, so a font relationship written
    into the document's rels would leave the id dangling, and python-docx
    would not notice.
    """
    run("embed", str(doc), fonts=licensed_fonts)
    pkg = package(doc)
    assert "word/_rels/fontTable.xml.rels" in pkg["names"]
    from lxml import etree
    document_rels = etree.fromstring(
        pkg["blobs"]["word/_rels/document.xml.rels"])
    for rel in document_rels.findall(f"{{{REL_NS}}}Relationship"):
        assert not rel.get("Target", "").endswith(".odttf")


def test_the_parts_are_typed_and_named_the_way_word_names_them(
        doc, licensed_fonts):
    run("embed", str(doc), "--family", FAMILY, fonts=licensed_fonts)
    pkg = package(doc)
    parts = sorted(n for n in pkg["names"] if n.endswith(".odttf"))
    assert parts == ["word/fonts/font1.odttf", "word/fonts/font2.odttf"]
    assert pkg["defaults"].get("odttf") == CT_OBFUSCATED


def test_regular_and_bold_land_in_separate_parts_with_separate_keys(
        doc, licensed_fonts):
    """Word tells the two apart by the part, and each part has its key."""
    report = run("embed", str(doc), fonts=licensed_fonts)
    slots = [r for r in report["embedded"]
             if r["family"] == FAMILY][0]["slots"]
    assert set(slots) == {"regular", "bold"}
    assert slots["regular"]["part"] != slots["bold"]["part"]
    entries = embedded_entries(package(doc))[FAMILY]
    assert entries["regular"][0] != entries["bold"][0]
    assert entries["regular"][1] != entries["bold"][1]


def test_the_embed_flags_are_set_in_the_settings(doc, licensed_fonts):
    run("embed", str(doc), fonts=licensed_fonts)
    settings = package(doc)["settings"]
    flag = settings.find(f"{{{W_NS}}}embedTrueTypeFonts")
    assert flag is not None
    assert flag.get(f"{{{W_NS}}}val") not in ("0", "false")
    # Whole faces are embedded, not subsets. An on/off element present
    # without a w:val means true, so this one needs the attribute.
    subset = settings.find(f"{{{W_NS}}}saveSubsetFonts")
    assert subset is not None and subset.get(f"{{{W_NS}}}val") == "false"
    # CT_Settings is a fixed sequence and these two sit early in it.
    order = [child.tag.split("}")[-1] for child in settings]
    for late in ("defaultTabStop", "compat", "rsids", "themeFontLang"):
        if late in order:
            assert order.index("embedTrueTypeFonts") < order.index(late)
            assert order.index("saveSubsetFonts") < order.index(late)
    assert order.index("embedTrueTypeFonts") < order.index("saveSubsetFonts")


def test_embedding_is_idempotent(doc, licensed_fonts):
    """A second pass adds no bytes and mints no new key.

    A fresh w:fontKey on the second run would be the quiet failure: the
    stored bytes were obfuscated for the old one, so the part would stop
    de-obfuscating to a font while every other check still passed.
    """
    run("embed", str(doc), fonts=licensed_fonts)
    first = package(doc)
    first_entries = embedded_entries(first)

    second_report = run("embed", str(doc), fonts=licensed_fonts)
    assert second_report["bytes_added"] == 0
    assert second_report["size_delta"] == 0
    second = package(doc)

    assert second["names"] == first["names"]
    assert embedded_entries(second) == first_entries
    for name in ("word/fontTable.xml", "word/settings.xml",
                 "[Content_Types].xml", "word/_rels/fontTable.xml.rels"):
        assert second["blobs"][name] == first["blobs"][name], name
    for name in first["names"]:
        if name.endswith(".odttf"):
            assert second["blobs"][name] == first["blobs"][name]
    assert_consistent(doc)


def test_report_on_an_embedded_document_proposes_nothing(doc,
                                                         licensed_fonts):
    """A dry run has to see the existing embedding, not plan a second."""
    run("embed", str(doc), "--family", FAMILY, fonts=licensed_fonts)
    report = run("report", str(doc), "--family", FAMILY, fonts=licensed_fonts)
    assert report["bytes_added"] == 0
    slots = [r for r in report["embedded"]
             if r["family"] == FAMILY][0]["slots"]
    assert slots and not any(slot["added"] for slot in slots.values())


def test_a_document_with_nothing_to_embed_grows_no_font_parts(
        doc, unlicensed_fonts):
    before = package(doc)
    run("embed", str(doc), "--family", FAMILY, fonts=unlicensed_fonts)
    after = package(doc)
    assert not [n for n in after["names"] if n.endswith(".odttf")]
    assert "odttf" not in after["defaults"]
    assert after["names"] == before["names"]
    assert reopens(doc)


def test_output_flag_leaves_the_source_untouched(doc, licensed_fonts,
                                                 tmp_path):
    out = tmp_path / "shipped.docx"
    run("embed", str(doc), "-o", str(out), fonts=licensed_fonts)
    assert [n for n in package(out)["names"] if n.endswith(".odttf")]
    assert not [n for n in package(doc)["names"] if n.endswith(".odttf")]


def test_verify_catches_a_part_wired_to_the_wrong_key(doc, licensed_fonts,
                                                     tmp_path):
    """The check that a round trip through this script cannot make.

    XOR is its own inverse, so a wrong key still round trips. What it
    cannot do is produce a real sfnt tag, which is what verify looks at.
    """
    run("embed", str(doc), fonts=licensed_fonts)
    broken = tmp_path / "broken.docx"
    pkg = package(doc)
    from lxml import etree
    table = pkg["table"]
    element = table.find(f"{{{W_NS}}}font/{{{W_NS}}}embedRegular")
    if element is None:
        for font in table.findall(f"{{{W_NS}}}font"):
            element = font.find(f"{{{W_NS}}}embedRegular")
            if element is not None:
                break
    assert element is not None
    element.set(f"{{{W_NS}}}fontKey",
                "{FFFFFFFF-FFFF-FFFF-FFFF-FFFFFFFFFFFF}")
    with zipfile.ZipFile(broken, "w", zipfile.ZIP_DEFLATED) as z:
        for name in pkg["names"]:
            data = etree.tostring(table, xml_declaration=True,
                                  encoding="UTF-8", standalone=True) \
                if name == "word/fontTable.xml" else pkg["blobs"][name]
            z.writestr(name, data)
    result = _invoke(("verify", str(broken)), None)
    assert result.returncode == 1
    problems = json.loads(result.stdout)["problems"]
    assert any("de-obfuscate" in line for line in problems), problems


# ---------------------------------------------------------------------------
# Real faces on this machine
# ---------------------------------------------------------------------------


def _family_is_findable(family: str) -> bool:
    sys.path.insert(0, str(CLI.parent))
    try:
        import docx_embed_fonts
    except SystemExit:
        return False
    except ImportError:
        return False
    finally:
        sys.path.pop(0)
    faces, _ = docx_embed_fonts.find_faces(family)
    return bool(faces)


@pytest.mark.skipif(not _family_is_findable(ARABIC_FAMILY),
                    reason=f"{ARABIC_FAMILY} is not installed as .ttf")
def test_an_arabic_report_embeds_its_arabic_face(tmp_path):
    """The case the feature exists for.

    Without the embedded face this report renders as whatever the
    reader's machine substitutes for the family name, or as tofu.
    """
    arabic = _with_family(tmp_path, "arabic.docx", ARABIC_FAMILY, ARABIC_TEXT)
    report = run("embed", str(arabic))
    assert ARABIC_FAMILY in report["families_used"]
    record = [r for r in report["embedded"] if r["family"] == ARABIC_FAMILY]
    assert record, report["skipped"]
    assert "regular" in record[0]["slots"]
    assert_consistent(arabic)
    assert ARABIC_TEXT in reopens(arabic)
    assert run("verify", str(arabic))["problems"] == []


@pytest.mark.skipif(shutil.which("soffice") is None
                    or not _family_is_findable(ARABIC_FAMILY),
                    reason="LibreOffice or the Arabic face is missing")
def test_the_embedded_report_still_converts_to_pdf(tmp_path):
    """A file Word would offer to repair is a failure.

    There is no Word here, so this is what stands in for it: an
    independent implementation still parses the package and still lays
    the document out. On its own a correct PDF only says the package is
    intact; the test below is the one that shows the embedded bytes were
    read.
    """
    arabic = _with_family(tmp_path, "arabic.docx", ARABIC_FAMILY, ARABIC_TEXT)
    run("embed", str(arabic))
    profile = tmp_path / "loprofile"
    result = _soffice(["--convert-to", "pdf", "--outdir", str(tmp_path),
                       str(arabic)], profile)
    if result is None:
        pytest.skip("LibreOffice would not run here")
    pdf = tmp_path / "arabic.pdf"
    assert pdf.is_file() and pdf.stat().st_size > 1000
    assert pdf.read_bytes()[:5] == b"%PDF-"


def _kufi_ttf() -> Path | None:
    """An Arabic .ttf that is not what LibreOffice substitutes.

    The probe below needs a face distinguishable from the fallback, so
    that a PDF built from the embedded bytes cannot be confused with a
    PDF built from a substitution.
    """
    try:
        out = subprocess.run(["fc-list", "--format", "%{file}\n"],
                             capture_output=True, text=True, timeout=20)
    except (OSError, subprocess.SubprocessError):
        return None
    for line in out.stdout.splitlines():
        if "notokufiarabic" in line.lower().replace(" ", "") \
                and line.lower().endswith(".ttf"):
            return Path(line)
    return None


@pytest.mark.skipif(shutil.which("soffice") is None
                    or shutil.which("pdffonts") is None,
                    reason="LibreOffice or pdffonts is not installed")
def test_another_implementation_renders_from_the_embedded_bytes(tmp_path):
    """The strongest check available without Word.

    The document asks for a family no machine has installed, so there is
    nothing to resolve the name against. Converted without the embedding
    LibreOffice substitutes its Arabic fallback; converted with it, the
    PDF carries the face whose bytes were embedded. That difference is
    only reachable by de-obfuscating the part with the key in the font
    table, so it exercises the content type, the relationship from
    fontTable.xml, the key and the settings flag at once.
    """
    source = _kufi_ttf()
    if source is None:
        pytest.skip("no distinctive Arabic .ttf to probe with")
    probe = "HermesLOProbe"
    fonts = tmp_path / "probe-fonts"
    (fonts / probe).mkdir(parents=True)
    shutil.copy(source, fonts / probe / f"{probe}-Regular.ttf")
    (fonts / probe / "OFL.txt").write_text("OFL 1.1\n", encoding="utf-8")

    plain = _with_family(tmp_path, "plain.docx", probe, ARABIC_TEXT)
    embedded = _with_family(tmp_path, "embedded.docx", probe, ARABIC_TEXT)
    report = run("embed", str(embedded), "--family", probe, fonts=fonts)
    assert [r["family"] for r in report["embedded"]] == [probe]

    profile = tmp_path / "loprofile"
    faces = {}
    for path in (plain, embedded):
        if _soffice(["--convert-to", "pdf", "--outdir", str(tmp_path),
                     str(path)], profile) is None:
            pytest.skip("LibreOffice would not run here")
        pdf = tmp_path / (path.stem + ".pdf")
        assert pdf.is_file() and pdf.read_bytes()[:5] == b"%PDF-"
        listing = subprocess.run(["pdffonts", str(pdf)],
                                 capture_output=True, text=True, timeout=60)
        faces[path.stem] = listing.stdout.lower()

    assert "kufi" not in faces["plain"], \
        "the probe family was resolvable, so this proves nothing"
    assert "kufi" in faces["embedded"], \
        "the embedded bytes were not used to render the document"


def test_a_longer_family_name_is_not_mistaken_for_this_one(
        tmp_path, ttf, doc):
    """IBM Plex Sans must not claim IBMPlexSansArabic-Regular.ttf."""
    root = tmp_path / "lookalike"
    (root / "other").mkdir(parents=True)
    shutil.copy(ttf, root / "other" / f"{FAMILY}Extended-Regular.ttf")
    (root / "other" / "OFL.txt").write_text("OFL 1.1\n", encoding="utf-8")
    report = run("embed", str(doc), fonts=root)
    assert all(r["family"] != FAMILY for r in report["embedded"])
    refused = [s for s in report["skipped"] if s["family"] == FAMILY]
    assert refused and refused[0]["reason"] == "no font file found"


@pytest.mark.skipif(not _family_is_findable("Cairo"),
                    reason="Cairo is not installed as .ttf")
def test_a_variable_font_fills_one_slot_and_says_so(doc):
    """One file covers the whole weight axis.

    Word renders an embedded face at its default instance, so pointing
    regular and bold at the same bytes would ship the same face twice
    under two labels for no gain.
    """
    report = run("embed", str(doc), "--family", "Cairo",
                 "--allow-unlicensed")
    record = [r for r in report["embedded"] if r["family"] == "Cairo"]
    if not record:
        pytest.skip("Cairo is installed as static faces on this machine")
    slots = record[0]["slots"]
    paths = [slot["file"] for slot in slots.values()]
    assert len(paths) == len(set(paths)), "a file was used for two slots"
    if "variable" in " ".join(record[0]["notes"]).lower():
        assert set(slots) == {"regular"}
    assert_consistent(doc)
