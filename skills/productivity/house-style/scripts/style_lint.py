#!/usr/bin/env python3
"""Read a finished deliverable back and say what gives it away.

Three kinds of finding, and the split is the whole point of the tool:

* **prose** catches the writing that reads as machine-made. Long dashes
  first, because that is the house rule most often broken and the one
  that lands in a file somebody else opens, then the vocabulary and the
  sentence rhythm that a reader feels before they can name.
* **design** catches the layout that reads as a template: stock Office
  fonts and the stock accent, eight type sizes on one slide, a wall of
  words nobody at the back of the room will read, a spreadsheet still
  wearing its grey grid.
* **geometry** catches the damage a reader sees first and no text check
  can: a shape off the slide or in the margin, text that needs more
  height than its box has, two shapes sitting on each other, type that
  fails contrast against what is actually behind it, and edges that are
  close to aligned without being aligned. Boxes and estimates only, so
  it stays as fast as the rest: no rendering and no LibreOffice.

Usage:
    python style_lint.py report.docx deck.pptx book.xlsx notes.md
    python style_lint.py deck.pptx --json
    python style_lint.py draft.md --only prose
    python style_lint.py deck.pptx --only geometry

Exit code is 1 when anything at error level is found, so the check can
gate a delivery. Warnings are budgets: they are worth a look and they do
not fail the run on their own.

A hit is a prompt to look, not a verdict. The word lists generalise badly
across industries, and a good writer can use any word on them well. The
one rule with no exception is the dash.
"""
from __future__ import annotations

import argparse
import json
import re
import statistics
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
from house_style import (  # noqa: E402
    LIMITS, contrast_ratio, estimate_lines, is_arabic,
    line_spacing_for, load_theme, luminance,
)

# ---------------------------------------------------------------------------
# Prose rules
# ---------------------------------------------------------------------------

# Written as escapes on purpose: this file is itself swept by the rule it
# enforces, and a literal dash here would make the sweep unusable.
LONG_DASHES = {
    "\u2014": "em dash", "\u2013": "en dash", "\u2012": "figure dash",
    "\u2015": "horizontal bar", "\u2e3a": "two-em dash",
    "\u2e3b": "three-em dash",
}

PROSE_PATTERNS = [
    ("llm_lexicon", re.compile(
        r"\b(delve[sd]?|delving|intricate|commendable|meticulous(ly)?|surpass(es|ed)?"
        r"|elevate[sd]?|foster(s|ed|ing)?|tapestry|realm|landscape of|pivotal"
        r"|resonate[sd]?|testament to|underscore[sd]?|showcasing|paramount"
        r"|unwavering)\b", re.I),
     "a word measured to spike in machine-written text"),
    ("corporate_lexicon", re.compile(
        r"\b(world.class|best.in.class|state.of.the.art|cutting.edge|seamless(ly)?"
        r"|holistic|bespoke|leverag(e|es|ing)|utili[sz]e[sd]?|empower(s|ed|ing)?"
        r"|unlock(s|ing)?|harness(es|ing)?|spearhead(s|ing)?)\b", re.I),
     "brochure vocabulary: say the thing instead"),
    ("intention_not_action", re.compile(
        r"\b(aims? to|seeks? to|strives? to|is committed to|are committed to"
        r"|is designed to|works? to)\b", re.I),
     "an intention cannot be checked; write what is done"),
    ("negative_parallelism", re.compile(
        r"\bnot (just|only|merely) [^.,;]{2,60}[,;]? but\b", re.I),
     "a reversal carrying no information"),
    ("stock_opener", re.compile(
        r"(\bIn today'?s [a-z\- ]{0,20}(world|landscape|environment|market|era)\b"
        r"|^\s*(Moreover|Furthermore|Additionally|In conclusion|It'?s worth noting)\b)",
        re.I | re.M),
     "an opener that announces instead of saying"),
    ("manufactured_suspense", re.compile(
        r"(Here'?s what sets .{2,40} apart|The result\?|But here'?s the thing)", re.I),
     "suspense manufactured where there is none"),
    ("pompous_copula", re.compile(
        r"\b(serves as|stands as|plays a (key|vital|pivotal|crucial) role"
        r"|positions itself as)\b", re.I),
     "plain 'is' was fine"),
    ("vague_attribution", re.compile(
        r"\b(industry experts|leading standards|international best practice"
        r"|studies show|research suggests)\b", re.I),
     "name the source or drop the claim"),
    # Arabic: the same faults, in the phrasing they actually arrive in.
    ("arabic_filler", re.compile(
        r"(في عالم اليوم|في ظل التطور|تجدر الإشارة|مما لا شك فيه"
        r"|يسعى إلى|تسعى إلى|نفخر بأن|نحرص على أن|بشكل فعال|بشكل كبير"
        r"|الجدير بالذكر|علاوة على ذلك)"),
     "حشو يقرأ كنص آلة: احذفه أو اكتب الفعل نفسه"),
    ("arabic_calque", re.compile(
        r"(يلعب دورا|تلعب دورا|في نهاية المطاف|من خلال الاستفادة من"
        r"|يعتبر واحدا من|تعتبر واحدة من)"),
     "ترجمة حرفية عن الإنجليزية: أعد صياغتها بالعربية"),
]

# Phrases a checker must not flag, so the fix for machine prose does not
# become its own machine prose. Kept here as documentation of intent.
NOT_A_TELL = (
    "a short page", "a repeated noun", "a sentence fragment",
    "a bare list with no descriptions", "a number with no adjective",
    "plain 'is'",
)

SENTENCE_SPLIT = re.compile(r"(?<=[.!?؟।])\s+|\n{2,}|؛\s*")
TRICOLON = re.compile(r"\b\w+, \w+,? (and|و)\s+\w+\b", re.I)


def _sentences(text: str) -> list[str]:
    return [s.strip() for s in SENTENCE_SPLIT.split(text) if s.strip()]


def _words(text: str) -> list[str]:
    return re.findall(r"[\w؀-ۿ']+", text)


def prose_findings(text: str, where: str) -> list[dict]:
    out: list[dict] = []
    for char, name in LONG_DASHES.items():
        count = text.count(char)
        if count:
            sample = next((line.strip() for line in text.splitlines()
                           if char in line), "")
            out.append(_f("error", "long_dash", where,
                          f"{count} x {name}; use a comma, a colon, or two "
                          f"sentences", sample[:120]))
    for rule, pattern, why in PROSE_PATTERNS:
        hits = pattern.findall(text)
        if hits:
            shown = [h if isinstance(h, str) else next((x for x in h if x), "")
                     for h in hits[:4]]
            out.append(_f("warn", rule, where,
                          f"{len(hits)} hit(s): {why}", ", ".join(shown)))

    sentences = _sentences(text)
    lengths = [len(_words(s)) for s in sentences if len(_words(s)) > 1]
    if len(lengths) >= 10:
        spread = statistics.pstdev(lengths)
        if spread < 6.0:
            out.append(_f("warn", "uniform_rhythm", where,
                          f"sentence length spread {spread:.1f} words, under 6: "
                          "put a five word sentence next to a twenty word one"))
    words = len(_words(text))
    tricolons = len(TRICOLON.findall(text))
    if words >= 200 and tricolons / (words / 400) > 1:
        out.append(_f("warn", "tricolon_reflex", where,
                      f"{tricolons} three item lists per {words} words"))
    return out


def sibling_openings(items: list[str], where: str) -> list[dict]:
    """Cards, bullets and list items that all start the same way."""
    starts: dict[str, int] = {}
    for item in items:
        w = _words(item)[:2]
        if len(w) == 2:
            key = " ".join(x.lower() for x in w)
            starts[key] = starts.get(key, 0) + 1
    repeats = {k: v for k, v in starts.items() if v >= 3}
    if repeats:
        worst = max(repeats.items(), key=lambda kv: kv[1])
        return [_f("warn", "repeated_opening", where,
                   f"{worst[1]} items open with \"{worst[0]}\"")]
    return []


def _f(level, rule, where, detail, sample=""):
    finding = {"level": level, "rule": rule, "where": where, "detail": detail}
    if sample:
        finding["sample"] = sample
    return finding


# ---------------------------------------------------------------------------
# Per format readers. Each returns (text, findings).
# ---------------------------------------------------------------------------

STOCK_FONTS = {"calibri", "calibri light", "aptos", "aptos display",
               "times new roman"}
STOCK_ACCENTS = {"4472C4", "ED7D31", "A5A5A5", "FFC000", "5B9BD5", "70AD47"}


def read_text_file(path: Path):
    return path.read_text(encoding="utf-8", errors="replace"), []


def _docx_loud_rules(table) -> bool:
    """A ruled table is fine. A table ruled in near black is not.

    The house grid is a light grey hairline, which a reader takes as
    structure. Word's own "Table Grid" rules every cell in the automatic
    color, which lands close to black and is what makes a pasted table
    shout across a page. So this reads the color and the weight rather
    than asking whether rules exist at all.
    """
    from docx.oxml.ns import qn

    borders = table._tbl.tblPr.find(qn("w:tblBorders"))
    if borders is None:
        style = getattr(table.style, "name", "") or ""
        return "grid" in style.lower()
    for edge in ("insideV", "insideH", "left", "right", "top", "bottom"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None or node.get(qn("w:val")) in (None, "none", "nil"):
            continue
        color = (node.get(qn("w:color")) or "auto").lstrip("#")
        if color.lower() == "auto":
            return True
        try:
            if luminance(color) < 0.45:
                return True
        except ValueError:
            return True
        if int(node.get(qn("w:sz")) or 0) > 12:
            return True
    return False


def read_docx(path: Path):
    from docx import Document
    from docx.shared import Pt

    doc = Document(str(path))
    chunks, findings, sizes, fonts = [], [], set(), set()
    for para in doc.paragraphs:
        chunks.append(para.text)
        words = len(_words(para.text))
        if words > LIMITS["doc_paragraph_words"]:
            findings.append(_f("warn", "long_paragraph", _where(path),
                               f"a paragraph of {words} words; split it"))
        for run in para.runs:
            if run.font.size:
                sizes.add(round(run.font.size.pt, 1))
            if run.font.name:
                fonts.add(run.font.name.lower())
    for style_name in ("Normal", "Heading 1", "Heading 2", "Body Text"):
        try:
            style = doc.styles[style_name]
        except KeyError:
            continue
        if style.font.name:
            fonts.add(style.font.name.lower())
        if style.font.size:
            sizes.add(round(style.font.size.pt, 1))
    for table in doc.tables:
        if _docx_loud_rules(table):
            findings.append(_f("warn", "loud_table_rules", _where(path),
                               "table ruled in near black; the house grid is a "
                               "light grey hairline that sits under the text"))
            break
        for row in table.rows:
            for cell in row.cells:
                chunks.append(cell.text)

    stock = fonts & STOCK_FONTS
    if stock:
        findings.append(_f("warn", "stock_font", _where(path),
                           f"stock office face in use: {', '.join(sorted(stock))}"))
    off_scale = _off_scale(sizes, "doc")
    if off_scale:
        findings.append(_f("warn", "off_scale_type", _where(path),
                           f"sizes outside the house scale: {sorted(off_scale)}"))
    elif len(sizes) > LIMITS["doc_font_sizes"] + 2:
        findings.append(_f("warn", "size_soup", _where(path),
                           f"{len(sizes)} distinct type sizes: "
                           f"{sorted(sizes)}"))
    # Measure: characters a line at the body size and the page width.
    section = doc.sections[0] if doc.sections else None
    if section is not None:
        width_pt = (section.page_width - section.left_margin
                    - section.right_margin) / 12700
        body = 11.0
        try:
            body = doc.styles["Normal"].font.size.pt or 11.0
        except (KeyError, AttributeError):
            pass
        chars = width_pt / (0.5 * body)
        low, high = LIMITS["doc_chars_per_line"]
        if not low <= chars <= high:
            findings.append(_f("warn", "measure", _where(path),
                               f"about {chars:.0f} characters a line; a "
                               f"typeset page sits between {low} and {high}"))
    return "\n".join(chunks), findings


def read_pptx(path: Path):
    from pptx import Presentation

    prs = Presentation(str(path))
    chunks, findings, sizes, fonts = [], [], set(), set()
    for n, slide in enumerate(prs.slides, 1):
        where = f"{_where(path)} slide {n}"
        words, bullets, visuals = 0, 0, 0
        items = []
        # A composed layout (cards, a timeline, a comparison) is many
        # positioned shapes, not a bullet list, and the shapes themselves
        # are the visual. Counting their paragraphs as bullets reported a
        # designed slide as a wall of text.
        composed = sum(1 for sh in slide.shapes
                       if not sh.is_placeholder and sh.has_text_frame) >= 3
        for shape in slide.shapes:
            if getattr(shape, "has_chart", False) or shape.shape_type == 13:
                visuals += 1
            if shape.has_table:
                visuals += 1
            if not shape.has_text_frame:
                continue
            for para in shape.text_frame.paragraphs:
                text = "".join(r.text for r in para.runs)
                if not text.strip():
                    continue
                chunks.append(text)
                items.append(text)
                words += len(_words(text))
                if (shape.is_placeholder
                        and shape.placeholder_format.idx != 0
                        and not composed):
                    bullets += 1
                    if len(_words(text)) > LIMITS["bullet_words"]:
                        findings.append(_f("warn", "long_bullet", where,
                                           f"{len(_words(text))} words in one "
                                           "bullet; target six"))
                for run in para.runs:
                    if run.font.size:
                        sizes.add(round(run.font.size.pt, 1))
                    if run.font.name:
                        fonts.add(run.font.name.lower())
                    try:
                        rgb = run.font.color.rgb
                    except (AttributeError, TypeError):
                        rgb = None
                    if rgb is not None and str(rgb) in STOCK_ACCENTS:
                        findings.append(_f("warn", "stock_accent", where,
                                           f"office theme accent {rgb} in text"))
        if words > LIMITS["slide_words_max"]:
            findings.append(_f("error", "slide_wall_of_text", where,
                               f"{words} words; split the slide, do not shrink "
                               f"the type (ceiling {LIMITS['slide_words_max']})"))
        elif words > LIMITS["slide_words"]:
            findings.append(_f("warn", "slide_heavy", where,
                               f"{words} words; target {LIMITS['slide_words']}"))
        if bullets > LIMITS["slide_bullets"]:
            findings.append(_f("warn", "bullet_pile", where,
                               f"{bullets} bullets; {LIMITS['slide_bullets']} "
                               "is the ceiling, and a grid usually reads better"))
        if composed:
            visuals += 1
        if visuals == 0 and words > 15:
            findings.append(_f("warn", "text_only_slide", where,
                               "no chart, table or image: text only slides are "
                               "the default failure of a generated deck"))
        findings.extend(sibling_openings(items, where))

    small = [s for s in sizes if s < 14]
    if small:
        findings.append(_f("warn", "type_too_small", _where(path),
                           f"type at {sorted(small)} pt; 18 is the floor for a "
                           "projected slide, 14 for one read as a file"))
    # A deck that uses eight roles of one scale is not size soup. A deck
    # carrying 17, 19 and 23 pt is, whatever the count says.
    off_scale = _off_scale(sizes, "deck")
    if off_scale:
        findings.append(_f("warn", "off_scale_type", _where(path),
                           f"sizes outside the house scale: {sorted(off_scale)}"))
    elif len(sizes) > LIMITS["deck_font_sizes"] + 2:
        findings.append(_f("warn", "size_soup", _where(path),
                           f"{len(sizes)} distinct type sizes: {sorted(sizes)}"))
    stock = fonts & STOCK_FONTS
    if stock:
        findings.append(_f("warn", "stock_font", _where(path),
                           f"stock office face in use: {', '.join(sorted(stock))}"))
    if len(fonts) > LIMITS["font_families"] + 1:
        findings.append(_f("warn", "font_zoo", _where(path),
                           f"{len(fonts)} families: {sorted(fonts)}"))
    return "\n".join(chunks), findings


def read_xlsx(path: Path):
    from openpyxl import load_workbook

    wb = load_workbook(str(path), data_only=False)
    chunks, findings, fonts = [], [], set()
    for ws in wb.worksheets:
        where = f"{_where(path)}[{ws.title}]"
        if ws.sheet_view.showGridLines:
            findings.append(_f("warn", "gridlines_on", where,
                               "the grey grid is the loudest thing on the "
                               "sheet; turn it off"))
        if ws.freeze_panes is None and ws.max_row > 12:
            findings.append(_f("warn", "no_freeze", where,
                               "a long sheet with no frozen header"))
        if ws.merged_cells.ranges:
            findings.append(_f("warn", "merged_cells", where,
                               f"{len(ws.merged_cells.ranges)} merged range(s): "
                               "they break sorting, filtering and formulas"))
        rules = getattr(ws.conditional_formatting, "_cf_rules", {})
        total = sum(len(v) for v in rules.values()) if rules else 0
        if total > LIMITS["conditional_formats_per_sheet"]:
            findings.append(_f("warn", "conditional_noise", where,
                               f"{total} conditional formatting rules"))
        plain_numbers = 0
        for row in ws.iter_rows():
            for cell in row:
                if cell.value is None:
                    continue
                if isinstance(cell.value, str):
                    chunks.append(cell.value)
                elif isinstance(cell.value, (int, float)) and \
                        cell.number_format == "General":
                    plain_numbers += 1
                if cell.font and cell.font.name:
                    fonts.add(cell.font.name.lower())
        if plain_numbers > 5:
            findings.append(_f("warn", "unformatted_numbers", where,
                               f"{plain_numbers} numbers still on General; give "
                               "each column a real number format"))
    if len(fonts) > 2:
        findings.append(_f("warn", "font_zoo", _where(path),
                           f"{len(fonts)} font families in one workbook"))
    return "\n".join(chunks), findings


def read_pdf(path: Path):
    import shutil
    import subprocess
    exe = shutil.which("pdftotext")
    if not exe:
        return "", [_f("warn", "no_extractor", _where(path),
                       "pdftotext is not installed; prose was not read")]
    try:
        out = subprocess.run([exe, str(path), "-"], capture_output=True,
                             text=True, timeout=120)
    except subprocess.TimeoutExpired:
        return "", [_f("warn", "no_extractor", _where(path),
                       "pdftotext did not finish within 120s; prose was not read")]
    if out.returncode != 0:
        detail = (out.stderr or "").strip().splitlines()
        reason = detail[-1] if detail else f"exit {out.returncode}"
        return "", [_f("warn", "no_extractor", _where(path),
                       f"pdftotext failed, prose was not read: {reason}")]
    return out.stdout, []


READERS = {
    ".docx": read_docx, ".pptx": read_pptx, ".xlsx": read_xlsx,
    ".pdf": read_pdf, ".md": read_text_file, ".txt": read_text_file,
    ".markdown": read_text_file, ".json": read_text_file,
}


def _off_scale(sizes, medium: str) -> set:
    """Sizes that belong to no role of the house scale for this medium."""
    theme = load_theme()
    scale = set(theme.deck.values() if medium == "deck" else theme.doc.values())
    # A stat or a cover number is allowed to go above the top of the
    # scale; it is the odd sizes in the middle that read as improvised.
    ceiling = max(scale)
    return {s for s in sizes if s not in scale and s < ceiling}


def _where(path: Path) -> str:
    """A name the reader can act on. Two SKILL.md files are not the same
    finding, and printing both as "SKILL.md" made that impossible to see."""
    try:
        return str(path.resolve().relative_to(Path.cwd()))
    except ValueError:
        return str(path)


# ---------------------------------------------------------------------------
# Geometry rules
# ---------------------------------------------------------------------------
#
# Everything above this line reads text and formatting, which is why text
# running past its box, a card sitting on top of its neighbour and pale
# type on a pale card all passed the lint. These rules read the boxes.
#
# No renderer is involved. Heights are estimated from the same line count
# the builder sizes its boxes with, so the two agree by construction, and
# a finding that rests on the estimate says so.

EMU_IN = 914400.0
# About 1.5 pt. Under this two boxes only touch, which is how a rule, a
# dot on a rule and a table border are drawn on purpose.
HAIRLINE_IN = 0.02
# A source line typeset to land on the bottom margin lands a hundredth of
# an inch either side of it. That is the margin working, not a defect.
MARGIN_TOL_IN = 0.03
# The line estimate is good to about five percent, and a frame is often
# sized to its estimate exactly, so only a real spill is reported.
OVERFLOW_SLACK = 1.25
# Edges this far apart in points read as a miss rather than a choice.
# Equal edges are the correct case and never fire.
NEAR_ALIGN_PT = (1.0, 4.0)
# 24 pt itself counts as large type, so it is scored against 3.0.
LARGE_TYPE_PT = 24.0
PICTURE_BACKDROP = "picture"


def _inches(value):
    return None if value is None else value / EMU_IN


def _shape_box(shape, slide):
    """Left, top, width, height in inches, or None when unresolvable.

    A placeholder can inherit its position from the layout, in which case
    python-pptx hands back None for all four. Every geometry rule touches
    a box, so resolve from the layout first and drop the shape if even
    that leaves a hole.
    """
    left, top, width, height = shape.left, shape.top, shape.width, shape.height
    if None in (left, top, width, height) and shape.is_placeholder:
        try:
            ph = slide.slide_layout.placeholders[shape.placeholder_format.idx]
        except (KeyError, AttributeError, ValueError):
            ph = None
        if ph is not None:
            left = ph.left if left is None else left
            top = ph.top if top is None else top
            width = ph.width if width is None else width
            height = ph.height if height is None else height
    if None in (left, top, width, height):
        return None
    return (_inches(left), _inches(top), _inches(width), _inches(height))


def _shape_fill(shape):
    """The hex of an opaque solid fill, "picture", or None.

    fill.type and fore_color.rgb both raise when the fill is inherited or
    carries a theme color rather than a literal one, which is the common
    case on a template slide.
    """
    if getattr(shape, "shape_type", None) == 13:
        return PICTURE_BACKDROP
    try:
        fill = shape.fill
        if fill.type != 1:               # MSO_FILL.SOLID
            return None
        return str(fill.fore_color.rgb)
    except (AttributeError, TypeError, ValueError, KeyError):
        return None


def _slide_background(slide, theme):
    for holder in (slide, slide.slide_layout, slide.slide_layout.slide_master):
        try:
            fill = holder.background.fill
            if fill.type == 1:
                return str(fill.fore_color.rgb)
        except (AttributeError, TypeError, ValueError, KeyError):
            continue
    return theme.color("paper")


def _layout_size(shape, slide):
    """The size a placeholder inherits, when the run does not carry one."""
    if not shape.is_placeholder:
        return None
    try:
        ph = slide.slide_layout.placeholders[shape.placeholder_format.idx]
    except (KeyError, AttributeError, ValueError):
        return None
    if not ph.has_text_frame:
        return None
    for para in ph.text_frame.paragraphs:
        if para.font.size:
            return para.font.size.pt
        for run in para.runs:
            if run.font.size:
                return run.font.size.pt
    return None


def _para_size(para, inherited):
    sizes = [r.font.size.pt for r in para.runs if r.font.size]
    if sizes:
        return max(sizes)
    if para.font.size:
        return para.font.size.pt
    return inherited


def _line_height_in(para, size_pt, text):
    """Height of one line in inches, honouring an exact leading."""
    from pptx.util import Length

    spacing = para.line_spacing
    if isinstance(spacing, Length):      # Length subclasses int, so test first
        return spacing / EMU_IN
    if isinstance(spacing, (int, float)):
        return size_pt * float(spacing) / 72.0
    return size_pt * line_spacing_for(text) / 72.0


def _frame_estimate(shape, slide):
    """(estimated height, usable height) in inches, or None to skip.

    Estimated the way the builder budgets: one call to estimate_lines per
    paragraph, times that paragraph's leading, plus its own spacing. A
    frame that grows to its text, does not wrap, or carries type of no
    resolvable size is not measurable here and is skipped rather than
    guessed at.
    """
    frame = shape.text_frame
    try:
        if frame.auto_size is not None and int(frame.auto_size) == 1:
            return None                  # MSO_AUTO_SIZE.SHAPE_TO_FIT_TEXT
    except (AttributeError, TypeError, ValueError):
        pass
    if frame.word_wrap is False:
        return None
    box = _shape_box(shape, slide)
    if box is None:
        return None
    _, _, width_in, height_in = box
    # Wrapping is measured across the box, the way the builder measures
    # it. Subtracting the left and right insets instead turned an
    # 81 character title into three lines where the renderer sets two,
    # and one extra line out of two is larger than any slack can cover.
    # The top and bottom insets are subtracted, because the text really
    # does start below them.
    pad_y = (_inches(frame.margin_top or 0) or 0.0) + \
            (_inches(frame.margin_bottom or 0) or 0.0)
    usable_h = height_in - pad_y
    if width_in <= 0 or usable_h <= 0:
        return None

    inherited = _layout_size(shape, slide)
    total = 0.0
    tallest_line = 0.0
    seen_text = False
    for para in frame.paragraphs:
        text = "".join(r.text for r in para.runs) or para.text or ""
        size = _para_size(para, inherited)
        if size is None:
            return None
        if text.strip():
            seen_text = True
        line_h = _line_height_in(para, size, text)
        tallest_line = max(tallest_line, line_h)
        total += estimate_lines(text, width_in, size) * line_h
        for gap in (para.space_before, para.space_after):
            if gap:
                total += gap / EMU_IN
    if not seen_text:
        return None
    return total, usable_h, tallest_line


def _slide_records(slide, theme):
    records = []
    for z, shape in enumerate(slide.shapes):
        box = _shape_box(shape, slide)
        if box is None:
            continue
        text = ""
        if shape.has_text_frame:
            text = shape.text_frame.text or ""
        records.append({
            "shape": shape, "z": z, "box": box, "text": text,
            "name": (shape.name or "").strip() or f"shape {z + 1}",
            "fill": _shape_fill(shape),
        })
    return records


def _contains(outer, inner, tol=MARGIN_TOL_IN):
    ol, ot, ow, oh = outer
    il, it, iw, ih = inner
    return (ol <= il + tol and ot <= it + tol
            and ol + ow >= il + iw - tol and ot + oh >= it + ih - tol)


def _full_bleed(box, slide_w, slide_h):
    _, _, w, h = box
    return w >= slide_w * 0.98 and h >= slide_h * 0.98


def _visible(record):
    """An empty unfilled text box paints nothing, so it cannot collide."""
    return bool(record["text"].strip()) or record["fill"] is not None


# 1. Off the slide, or inside the safe margin.

def _margin_findings(records, where, slide_w, slide_h, geometry):
    # The house margins are stated for the 13.333 x 7.5 in canvas. A 4:3
    # deck is a smaller sheet, not a deck with fatter margins, so scale
    # them to the canvas actually in hand before judging anything.
    scale_x = slide_w / geometry["width"]
    scale_y = slide_h / geometry["height"]
    margin_x = geometry["margin_x"] * scale_x
    margin_top = geometry["margin_top"] * scale_y
    margin_bottom = geometry["margin_bottom"] * scale_y
    findings = []
    for rec in records:
        left, top, width, height = rec["box"]
        past = {}
        if left < -MARGIN_TOL_IN:
            past["left"] = -left
        if top < -MARGIN_TOL_IN:
            past["top"] = -top
        if left + width > slide_w + MARGIN_TOL_IN:
            past["right"] = left + width - slide_w
        if top + height > slide_h + MARGIN_TOL_IN:
            past["bottom"] = top + height - slide_h
        if past:
            edge, over = max(past.items(), key=lambda kv: kv[1])
            findings.append(_f("warn", "off_slide", where,
                               f'"{rec["name"]}" runs {over:.2f} in past the '
                               f"{edge} edge of the slide"))
            continue
        # A deliberate full bleed fills the margin by design.
        if _full_bleed(rec["box"], slide_w, slide_h):
            continue
        into = {
            "left": margin_x - left,
            "top": margin_top - top,
            "right": (left + width) - (slide_w - margin_x),
            "bottom": (top + height) - (slide_h - margin_bottom),
        }
        bad = {k: v for k, v in into.items() if v > MARGIN_TOL_IN}
        if bad:
            edge, over = max(bad.items(), key=lambda kv: kv[1])
            findings.append(_f("warn", "out_of_margin", where,
                               f'"{rec["name"]}" crosses the {edge} safe '
                               f"margin by {over:.2f} in"))
    return findings


# 2. Text that needs more height than its box has.

def _overflow_findings(records, slide, where):
    findings = []
    for rec in records:
        shape = rec["shape"]
        if not shape.has_text_frame or not rec["text"].strip():
            continue
        estimate = _frame_estimate(shape, slide)
        if estimate is None:
            continue
        needed, usable, line_h = estimate
        # Two gates. The line count is a ceiling, so one character over a
        # wrap boundary buys a whole extra line: an Arabic cover title of
        # 31 characters in a 30 character measure was estimated at two
        # lines and set as one. A frame over by less than three quarters
        # of a line of its own type is inside that noise.
        if needed > usable * OVERFLOW_SLACK and \
                needed - usable > 0.75 * line_h:
            findings.append(_f(
                "error", "text_overflow", where,
                f'"{rec["name"]}" needs about {needed:.2f} in of height in a '
                f"box {usable:.2f} in tall, {needed / usable:.2f}x its room. "
                "Line count is an estimate, not a render, so the number is "
                "approximate; the spill is not.",
                rec["text"].strip().replace("\n", " ")[:120]))
    return findings


# 3. Two boxes that intersect.

def _overlap_findings(records, where):
    findings = []
    visible = [r for r in records if _visible(r)]
    for i, a in enumerate(visible):
        for b in visible[i + 1:]:
            al, at, aw, ah = a["box"]
            bl, bt, bw, bh = b["box"]
            if min(aw, ah) < HAIRLINE_IN or min(bw, bh) < HAIRLINE_IN:
                continue                 # a rule, a dot, a hairline divider
            over_w = min(al + aw, bl + bw) - max(al, bl)
            over_h = min(at + ah, bt + bh) - max(at, bt)
            if over_w <= HAIRLINE_IN or over_h <= HAIRLINE_IN:
                continue
            # A backdrop is not a collision: a card, a band or a full
            # bleed holds the shape it sits behind on purpose.
            if _backdrop_pair(a, b) or _backdrop_pair(b, a):
                continue
            findings.append(_f("warn", "shape_overlap", where,
                               f'"{a["name"]}" and "{b["name"]}" overlap by '
                               f"{over_w:.2f} x {over_h:.2f} in"))
    return findings


def _backdrop_pair(outer, inner):
    """True when outer is deliberately behind inner."""
    if outer["z"] > inner["z"]:
        return False
    if outer["fill"] is None:
        return False
    if outer["text"].strip():
        return False
    return _contains(outer["box"], inner["box"])


# 4. Contrast against whatever is actually behind the text.

def _backdrop_for(records, index, background):
    """Walk z order and keep the last opaque thing that holds this box."""
    box = records[index]["box"]
    backdrop = background
    for other in records[:index]:
        if other["fill"] is None:
            continue
        if _contains(other["box"], box):
            backdrop = other["fill"]
    # A shape's own fill is painted directly under its own text and wins
    # over everything below it. Without this last step, white text in a
    # filled callout was scored against the white slide behind the
    # callout and reported as unreadable.
    own = records[index]["fill"]
    return own if own is not None else backdrop


def _contrast_findings(records, slide, where, background, theme):
    findings = []
    seen = set()
    for i, rec in enumerate(records):
        shape = rec["shape"]
        if not shape.has_text_frame or not rec["text"].strip():
            continue
        backdrop = _backdrop_for(records, i, background)
        if backdrop == PICTURE_BACKDROP:
            continue                     # a photograph has no one luminance
        inherited = _layout_size(shape, slide)
        for para in shape.text_frame.paragraphs:
            for run in para.runs:
                if not run.text.strip():
                    continue
                try:
                    ink = str(run.font.color.rgb)
                except (AttributeError, TypeError, ValueError):
                    continue
                size = (run.font.size.pt if run.font.size
                        else _para_size(para, inherited))
                if size is None:
                    continue
                try:
                    ratio = contrast_ratio(ink, backdrop)
                except (ValueError, TypeError):
                    continue
                floor = 3.0 if size >= LARGE_TYPE_PT else 4.5
                if ratio >= floor:
                    continue
                key = (rec["name"], ink, backdrop)
                if key in seen:
                    continue
                seen.add(key)
                findings.append(_f(
                    "warn", "low_contrast_in_place", where,
                    f'"{rec["name"]}" sets {size:.0f} pt {ink} on {backdrop}: '
                    f"contrast {ratio:.2f}, floor {floor}",
                    run.text.strip()[:80]))
    return findings


# 5. Edges that are close but not equal.

def _alignment_findings(records, where):
    findings = []
    edges = {"left": [], "right": []}
    for rec in records:
        left, _, width, height = rec["box"]
        if width < HAIRLINE_IN or height < HAIRLINE_IN or not _visible(rec):
            continue
        edges["left"].append((left, rec["name"]))
        edges["right"].append((left + width, rec["name"]))
    for kind, items in edges.items():
        items.sort()
        for (a, first), (b, second) in zip(items, items[1:]):
            delta = (b - a) * 72
            if NEAR_ALIGN_PT[0] <= delta <= NEAR_ALIGN_PT[1]:
                findings.append(_f("warn", "near_alignment", where,
                                   f'"{first}" and "{second}" {kind} edges '
                                   f"differ by {delta:.1f} pt: make them "
                                   "equal or make the difference visible"))
    return findings


def geometry_pptx(path: Path) -> list[dict]:
    from pptx import Presentation

    prs = Presentation(str(path))
    theme = load_theme()
    geometry = theme.geometry
    slide_w = _inches(prs.slide_width)
    slide_h = _inches(prs.slide_height)
    findings: list[dict] = []
    for n, slide in enumerate(prs.slides, 1):
        where = f"{_where(path)} slide {n}"
        records = _slide_records(slide, theme)
        background = _slide_background(slide, theme)
        findings += _margin_findings(records, where, slide_w, slide_h,
                                     geometry)
        findings += _overflow_findings(records, slide, where)
        findings += _overlap_findings(records, where)
        findings += _contrast_findings(records, slide, where, background,
                                       theme)
        findings += _alignment_findings(records, where)
    return findings


def _docx_shading(element, qn) -> str | None:
    if element is None:
        return None
    shd = element.find(qn("w:shd"))
    if shd is None:
        return None
    fill = (shd.get(qn("w:fill")) or "").lstrip("#")
    if not fill or fill.lower() == "auto":
        return None
    return fill.upper()


def _docx_runs(paragraph, where, backdrop, body_size, findings, seen):
    for run in paragraph.runs:
        if not run.text.strip():
            continue
        # A theme colored run carries no literal rgb, and reading it
        # raises rather than returning None.
        try:
            rgb = run.font.color.rgb
        except (AttributeError, TypeError, ValueError):
            continue
        if rgb is None:
            continue
        ink = str(rgb)
        size = run.font.size.pt if run.font.size else body_size
        try:
            ratio = contrast_ratio(ink, backdrop)
        except (ValueError, TypeError):
            continue
        floor = 3.0 if size >= LARGE_TYPE_PT else 4.5
        if ratio >= floor:
            continue
        key = (where, ink, backdrop)
        if key in seen:
            continue
        seen.add(key)
        findings.append(_f("warn", "low_contrast_in_place", where,
                           f"{size:.0f} pt {ink} on {backdrop}: contrast "
                           f"{ratio:.2f}, floor {floor}",
                           run.text.strip()[:80]))


def geometry_docx(path: Path) -> list[dict]:
    """The geometry rules that survive a reflowing page.

    Word sets its own line breaks and pushes what does not fit onto the
    next page, so a paragraph cannot overflow its box, there is no z
    order to collide in and no positioned edges to misalign. What does
    carry over is the measure, which an oversized image or table breaks,
    and the contrast of ink on the shading actually behind it.
    """
    from docx import Document
    from docx.oxml.ns import qn

    doc = Document(str(path))
    theme = load_theme()
    paper = theme.color("paper")
    body_size = theme.doc["body"]
    try:
        body_size = doc.styles["Normal"].font.size.pt or body_size
    except (KeyError, AttributeError):
        pass
    findings: list[dict] = []
    seen: set = set()
    if not doc.sections:
        return findings
    section = doc.sections[0]
    column_in = _inches(section.page_width - section.left_margin
                        - section.right_margin)

    for n, image in enumerate(doc.inline_shapes, 1):
        width_in = _inches(image.width)
        if width_in is None:
            continue
        if width_in > column_in + MARGIN_TOL_IN:
            findings.append(_f("warn", "out_of_margin",
                               f"{_where(path)} image {n}",
                               f"image {n} is {width_in:.2f} in wide in a "
                               f"{column_in:.2f} in column, over by "
                               f"{width_in - column_in:.2f} in"))

    for n, table in enumerate(doc.tables, 1):
        width_in = _docx_table_width(table, qn)
        if width_in is not None and width_in > column_in + MARGIN_TOL_IN:
            findings.append(_f("warn", "out_of_margin",
                               f"{_where(path)} table {n}",
                               f"table {n} is {width_in:.2f} in wide in a "
                               f"{column_in:.2f} in column, over by "
                               f"{width_in - column_in:.2f} in"))

    for i, para in enumerate(doc.paragraphs, 1):
        if not para.text.strip():
            continue
        backdrop = _docx_shading(para._p.pPr, qn) or paper
        _docx_runs(para, f"{_where(path)} paragraph {i}", backdrop,
                   body_size, findings, seen)
    for n, table in enumerate(doc.tables, 1):
        for r, row in enumerate(table.rows, 1):
            for c, cell in enumerate(row.cells, 1):
                backdrop = _docx_shading(cell._tc.tcPr, qn) or paper
                for para in cell.paragraphs:
                    if not para.text.strip():
                        continue
                    _docx_runs(para,
                               f"{_where(path)} table {n} cell {r}.{c}",
                               backdrop, body_size, findings, seen)
    return findings


def _docx_table_width(table, qn):
    """Width in inches: the wider of the preferred width and the grid.

    A table carries both, and Word lays out to whichever is larger, so
    reading only one of them let a table with a 9 in grid and an auto
    preferred width pass as if it fitted.
    """
    widths = []
    tbl_pr = table._tbl.tblPr
    if tbl_pr is not None:
        for node in tbl_pr.findall(qn("w:tblW")):
            if node.get(qn("w:type")) != "dxa":
                continue
            try:
                widths.append(int(node.get(qn("w:w"))) / 1440.0)
            except (TypeError, ValueError):
                continue
    grid = table._tbl.find(qn("w:tblGrid"))
    if grid is not None:
        total = 0
        for col in grid.findall(qn("w:gridCol")):
            try:
                total += int(col.get(qn("w:w")))
            except (TypeError, ValueError):
                total = 0
                break
        if total:
            widths.append(total / 1440.0)
    return max(widths) if widths else None


GEOMETRY = {".pptx": geometry_pptx, ".docx": geometry_docx}


def geometry_findings(path: Path) -> list[dict]:
    reader = GEOMETRY.get(path.suffix.lower())
    return reader(path) if reader else []


def lint(path: Path, only: str = "all") -> list[dict]:
    reader = READERS.get(path.suffix.lower())
    if reader is None:
        return [_f("warn", "unsupported", _where(path),
                   f"no reader for {path.suffix}")]
    text, findings = reader(path)
    if only not in ("all", "design"):
        findings = []
    if only in ("all", "geometry"):
        findings += geometry_findings(path)
    if only in ("all", "prose"):
        findings += prose_findings(text, _where(path))
    return findings


def main(argv=None):
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    parser = argparse.ArgumentParser(
        description="Lint a deliverable for machine-written prose and "
                    "template design.")
    parser.add_argument("files", nargs="+")
    parser.add_argument("--json", action="store_true", help="JSON output")
    parser.add_argument("--only", choices=("all", "prose", "design", "geometry"),
                        default="all")
    parser.add_argument("--strict", action="store_true",
                        help="fail on warnings too")
    args = parser.parse_args(argv)

    findings: list[dict] = []
    for name in args.files:
        path = Path(name)
        if not path.exists():
            findings.append(_f("error", "missing", name, "no such file"))
            continue
        findings.extend(lint(path, args.only))

    errors = [f for f in findings if f["level"] == "error"]
    warns = [f for f in findings if f["level"] == "warn"]
    if args.json:
        print(json.dumps({"ok": not errors, "errors": len(errors),
                          "warnings": len(warns), "findings": findings},
                         indent=2, ensure_ascii=False))
    else:
        for f in findings:
            mark = "FAIL" if f["level"] == "error" else "warn"
            line = f"{mark:>4}  {f['where']}: {f['rule']} - {f['detail']}"
            if f.get("sample"):
                line += f"\n        {f['sample']}"
            print(line)
        print(f"\n{len(errors)} error(s), {len(warns)} warning(s)")
    return 1 if errors or (args.strict and warns) else 0


if __name__ == "__main__":
    sys.exit(main())
