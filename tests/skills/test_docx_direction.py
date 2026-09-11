"""Direction in a document that carries two scripts.

A paragraph has one base direction. Its runs do not. Getting that wrong
is not a styling complaint, it changes what the sentence says: an Arabic
line with no base direction puts its full stop at the start, and a Latin
brand name marked right to left inside an Arabic sentence comes out
reversed.

These tests pin both halves, and they pin repair as separate from
restyling, because a document somebody else wrote gets its direction
fixed and nothing else touched.
"""

from __future__ import annotations

import json
import subprocess
import sys
import zipfile
from pathlib import Path

import pytest

REPO = Path(__file__).resolve().parents[2]
SCRIPTS = REPO / "skills" / "productivity" / "docx" / "scripts"
EDIT = SCRIPTS / "docx_edit.py"

AR = "البند الجزائي يحتاج مراجعة قبل التوقيع."
EN = "The notice period is now sixty days for either party."


@pytest.fixture(scope="module")
def common():
    if str(SCRIPTS) not in sys.path:
        sys.path.insert(0, str(SCRIPTS))
    import docx_common  # noqa: PLC0415

    return docx_common


def qn(tag):
    from docx.oxml.ns import qn as _qn

    return _qn(tag)


def run_rtl_flags(para) -> list[bool]:
    out = []
    for run in para.runs:
        rpr = run._r.find(qn("w:rPr"))
        out.append(rpr is not None and rpr.find(qn("w:rtl")) is not None)
    return out


def para_is_rtl(para) -> bool:
    ppr = para._p.find(qn("w:pPr"))
    return ppr is not None and ppr.find(qn("w:bidi")) is not None


def cli(*args) -> dict:
    out = subprocess.run([sys.executable, str(EDIT), *args],
                         capture_output=True, text=True)
    assert out.returncode == 0, out.stderr
    return json.loads(out.stdout)


def test_an_arabic_sentence_keeps_its_latin_run_left_to_right(common):
    """The case that reverses a brand name if you mark every run."""
    pytest.importorskip("docx")
    from docx import Document

    doc = Document()
    para = doc.add_paragraph()
    para.add_run("تم توريد ")
    para.add_run("IBM Plex Sans Arabic")
    para.add_run(" بتاريخ ")
    para.add_run("2026")
    common.set_paragraph_rtl(para, True)

    assert para_is_rtl(para), "the paragraph's base direction is Arabic"
    assert run_rtl_flags(para) == [True, False, True, True], \
        "the Latin run must not be marked, the digits follow the paragraph"


def test_an_english_sentence_marks_only_the_arabic_it_quotes(common):
    pytest.importorskip("docx")
    from docx import Document

    doc = Document()
    para = doc.add_paragraph()
    para.add_run("The clause is titled ")
    para.add_run("البند الجزائي")
    para.add_run(" in the Arabic copy.")
    common.mark_runs_by_script(para, base_rtl=False)

    assert not para_is_rtl(para), "an English sentence stays left to right"
    assert run_rtl_flags(para) == [False, True, False]


def test_the_auto_pass_separates_arabic_paragraphs_from_mixed_ones(common):
    pytest.importorskip("docx")
    from docx import Document

    doc = Document()
    arabic = doc.add_paragraph(AR)
    english = doc.add_paragraph(EN)
    mixed = doc.add_paragraph()
    mixed.add_run("The clause reads ")
    mixed.add_run("البند الجزائي")
    mixed.add_run(" in Arabic.")

    counts = common.apply_rtl(doc, "auto")

    assert para_is_rtl(arabic)
    assert not para_is_rtl(english)
    assert not para_is_rtl(mixed), "mostly Latin keeps its own base direction"
    assert counts["paragraphs_rtl"] == 1
    assert counts["mixed_paragraphs"] == 1
    assert run_rtl_flags(mixed) == [False, True, False]


def test_repair_changes_direction_and_nothing_else(tmp_path):
    """A file somebody else wrote: fix the direction, touch nothing more."""
    pytest.importorskip("docx")
    from docx import Document

    doc = Document()
    doc.add_heading("الملاحظات العربية", 1)
    doc.add_paragraph(AR)
    doc.add_paragraph(EN)
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "المنتج"
    table.cell(0, 1).text = "الهامش"
    table.cell(1, 0).text = "المقابض"
    table.cell(1, 1).text = "41%"
    source = tmp_path / "client.docx"
    doc.save(str(source))

    def text_of(path):
        d = Document(str(path))
        rows = [c.text for t in d.tables for r in t.rows for c in r.cells]
        return [p.text for p in d.paragraphs] + rows

    def styles_of(path):
        d = Document(str(path))
        return [(p.style.name, p.runs[0].font.size if p.runs else None)
                for p in d.paragraphs]

    before_text, before_styles = text_of(source), styles_of(source)
    out = tmp_path / "fixed.docx"
    report = cli("direction", str(source), "-o", str(out))

    assert report["paragraphs_rtl"] >= 2
    assert text_of(out) == before_text, "repair rewrote the text"
    assert styles_of(out) == before_styles, "repair restyled the document"

    fixed = Document(str(out))
    assert para_is_rtl(fixed.paragraphs[0]), "the Arabic heading was flipped"
    assert not para_is_rtl(fixed.paragraphs[2]), "the English line was not"
    assert fixed.tables[0]._tbl.tblPr.find(qn("w:bidiVisual")) is not None, \
        "an Arabic table reads from the right"


def test_repair_is_idempotent(tmp_path):
    pytest.importorskip("docx")
    from docx import Document

    doc = Document()
    doc.add_paragraph(AR)
    doc.add_paragraph(EN)
    source = tmp_path / "a.docx"
    doc.save(str(source))

    once = tmp_path / "once.docx"
    twice = tmp_path / "twice.docx"
    cli("direction", str(source), "-o", str(once))
    cli("direction", str(once), "-o", str(twice))

    def document_xml(path):
        with zipfile.ZipFile(path) as zf:
            return zf.read("word/document.xml")

    assert document_xml(once) == document_xml(twice), \
        "running the repair twice must not keep adding marks"


def test_forcing_the_whole_document_right_to_left(tmp_path):
    pytest.importorskip("docx")
    from docx import Document

    doc = Document()
    doc.add_paragraph(EN)
    source = tmp_path / "en.docx"
    doc.save(str(source))
    out = tmp_path / "forced.docx"

    report = cli("direction", str(source), "--mode", "on", "-o", str(out))
    assert report["mode"] == "on"
    assert para_is_rtl(Document(str(out)).paragraphs[0])

    off = tmp_path / "untouched.docx"
    report = cli("direction", str(source), "--mode", "off", "-o", str(off))
    assert report["paragraphs_rtl"] == 0
    assert not para_is_rtl(Document(str(off)).paragraphs[0])


def test_a_run_of_digits_alone_follows_its_paragraph(common):
    """A version number in an Arabic line is not a direction of its own."""
    pytest.importorskip("docx")
    from docx import Document

    doc = Document()
    para = doc.add_paragraph()
    para.add_run("النسخة ")
    para.add_run("2.4.1")
    common.set_paragraph_rtl(para, True)
    assert run_rtl_flags(para) == [True, True]

    latin = doc.add_paragraph()
    latin.add_run("Version ")
    latin.add_run("2.4.1")
    common.mark_runs_by_script(latin, base_rtl=False)
    assert run_rtl_flags(latin) == [False, False]
