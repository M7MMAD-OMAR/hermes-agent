"""Contracts for the docx comment conversation: threads, replies, resolve.

The docx skill is how Hermes holds a review conversation inside a Word
document, so the tests that matter are the ones that keep a thread a
thread: a reply must hang off the thread root, the resolved mark must
cover the whole thread, deleting a thread must leave no dangling anchor,
and a file that never had comments must grow the parts it needs without
becoming a file Word offers to repair.
"""

from __future__ import annotations

import importlib.util
import json
import shutil
import subprocess
import sys
import zipfile
from pathlib import Path

import pytest

REPO = Path(__file__).resolve().parents[2]
SCRIPTS = REPO / "skills" / "productivity" / "docx" / "scripts"

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
W14 = "http://schemas.microsoft.com/office/word/2010/wordml"
W15 = "http://schemas.microsoft.com/office/word/2012/wordml"

pytest.importorskip("docx", reason="python-docx is not installed")


def _load(name: str, path: Path):
    spec = importlib.util.spec_from_file_location(name, path)
    module = importlib.util.module_from_spec(spec)
    sys.modules[name] = module
    spec.loader.exec_module(module)
    return module


@pytest.fixture(scope="module")
def dc():
    if str(SCRIPTS) not in sys.path:
        sys.path.insert(0, str(SCRIPTS))
    return _load("docx_comments_under_test", SCRIPTS / "docx_comments.py")


@pytest.fixture(scope="module")
def validate():
    if str(SCRIPTS) not in sys.path:
        sys.path.insert(0, str(SCRIPTS))
    return _load("docx_validate_under_test", SCRIPTS / "docx_validate.py")


def _new_doc(path: Path, text: str = "Q3 revenue rose sharply this year."):
    from docx import Document

    doc = Document()
    doc.add_paragraph(text)
    doc.save(str(path))
    return path


def _commented(dc, path: Path, target: str = "Q3 revenue",
               text: str = "Needs a source", author: str = "Reviewer"):
    """A saved document carrying one comment on `target`."""
    from docx import Document

    doc = Document(str(path))
    _para, runs = dc.find_anchor_runs(doc, target)
    assert runs, f"anchor not found: {target}"
    cid = dc.add_comment_xml(doc, runs, text, author, "RV")
    dc.register_comment(doc, cid)
    dc.ensure_comment_styles(doc)
    doc.save(str(path))
    return cid


def _reopen(path: Path):
    from docx import Document

    return Document(str(path))


def _part(path: Path, name: str) -> str:
    with zipfile.ZipFile(str(path)) as zf:
        return zf.read(name).decode("utf-8")


def _names(path: Path) -> list:
    with zipfile.ZipFile(str(path)) as zf:
        return zf.namelist()


# ---------------------------------------------------------------------------
# Threading
# ---------------------------------------------------------------------------


def test_reply_carries_the_thread_root_as_para_id_parent(dc, tmp_path):
    """Both replies point at the root, not at each other.

    Word threads are flat, so a reply to a reply belongs to the same root.
    A chain would render as a broken conversation.
    """
    path = _new_doc(tmp_path / "reply.docx")
    root_id = _commented(dc, path)

    doc = _reopen(path)
    first = dc.reply_to_comment(doc, root_id, "Added the source", "Hermes")
    doc.save(str(path))

    doc = _reopen(path)
    second = dc.reply_to_comment(doc, first["comment_id"], "Checked", "QA")
    doc.save(str(path))

    assert first["parent_id"] == root_id
    assert second["parent_id"] == root_id
    assert first["para_id_parent"] == second["para_id_parent"]

    threads = dc.list_threads(_reopen(path))
    assert len(threads) == 1
    assert threads[0]["id"] == root_id
    assert [r["id"] for r in threads[0]["replies"]] == [
        first["comment_id"], second["comment_id"]]

    extended = _part(path, "word/commentsExtended.xml")
    parent = threads[0]["para_id"]
    assert extended.count(f'w15:paraIdParent="{parent}"') == 2


def test_reply_anchors_to_the_same_document_range(dc, tmp_path):
    path = _new_doc(tmp_path / "anchor.docx")
    root_id = _commented(dc, path)
    doc = _reopen(path)
    reply = dc.reply_to_comment(doc, root_id, "Agreed", "Hermes")
    doc.save(str(path))

    records = {c["id"]: c for c in dc.list_comments(_reopen(path))}
    assert records[reply["comment_id"]]["anchored_text"] == "Q3 revenue"
    assert records[root_id]["anchored_text"] == "Q3 revenue"


def test_flat_listing_still_reports_the_original_fields(dc, tmp_path):
    """Callers that read the flat shape keep working; keys were added only."""
    path = _new_doc(tmp_path / "flat.docx")
    _commented(dc, path)
    record = dc.list_comments(_reopen(path))[0]
    for key in ("id", "author", "initials", "date", "text", "anchored_text"):
        assert key in record
    assert record["author"] == "Reviewer"
    assert record["text"] == "Needs a source"


# ---------------------------------------------------------------------------
# Resolve and reopen
# ---------------------------------------------------------------------------


def test_thread_resolves_and_reopens_as_a_whole(dc, tmp_path):
    path = _new_doc(tmp_path / "resolve.docx")
    root_id = _commented(dc, path)
    doc = _reopen(path)
    dc.reply_to_comment(doc, root_id, "Added the source", "Hermes")
    doc.save(str(path))

    doc = _reopen(path)
    members = dc.set_thread_resolved(doc, root_id, True)
    doc.save(str(path))
    assert len(members) == 2
    assert all(c["resolved"] for c in dc.list_comments(_reopen(path)))

    # A reply to a resolved thread keeps the thread's state consistent.
    doc = _reopen(path)
    dc.reply_to_comment(doc, root_id, "One more", "QA")
    doc.save(str(path))
    assert all(c["resolved"] for c in dc.list_comments(_reopen(path)))

    doc = _reopen(path)
    dc.set_thread_resolved(doc, root_id, False)
    doc.save(str(path))
    comments = dc.list_comments(_reopen(path))
    assert len(comments) == 3
    assert not any(c["resolved"] for c in comments)
    assert 'w15:done="1"' not in _part(path, "word/commentsExtended.xml")


def test_resolving_from_a_reply_resolves_the_root(dc, tmp_path):
    path = _new_doc(tmp_path / "from-reply.docx")
    root_id = _commented(dc, path)
    doc = _reopen(path)
    reply = dc.reply_to_comment(doc, root_id, "Added the source", "Hermes")
    doc.save(str(path))

    doc = _reopen(path)
    members = dc.set_thread_resolved(doc, reply["comment_id"], True)
    doc.save(str(path))
    assert members[0] == root_id
    assert all(c["resolved"] for c in dc.list_comments(_reopen(path)))


# ---------------------------------------------------------------------------
# Deleting
# ---------------------------------------------------------------------------


def test_deleting_a_thread_leaves_no_dangling_anchor(dc, validate, tmp_path):
    path = _new_doc(tmp_path / "delete.docx")
    root_id = _commented(dc, path)
    doc = _reopen(path)
    dc.reply_to_comment(doc, root_id, "Added the source", "Hermes")
    doc.save(str(path))

    doc = _reopen(path)
    removed = dc.delete_thread(doc, root_id)
    doc.save(str(path))
    assert len(removed) == 2

    assert dc.list_comments(_reopen(path)) == []
    document = _part(path, "word/document.xml")
    assert "commentRangeStart" not in document
    assert "commentRangeEnd" not in document
    assert "commentReference" not in document
    assert "w15:commentEx " not in _part(path, "word/commentsExtended.xml")
    assert "w16cid:commentId " not in _part(path, "word/commentsIds.xml")
    assert validate.validate(str(path))["ok"] is True


def test_deleting_one_reply_keeps_the_rest_of_the_thread(dc, tmp_path):
    path = _new_doc(tmp_path / "delete-one.docx")
    root_id = _commented(dc, path)
    doc = _reopen(path)
    reply = dc.reply_to_comment(doc, root_id, "Added the source", "Hermes")
    doc.save(str(path))

    doc = _reopen(path)
    assert dc.delete_comment(doc, reply["comment_id"]) is True
    doc.save(str(path))

    threads = dc.list_threads(_reopen(path))
    assert len(threads) == 1
    assert threads[0]["replies"] == []
    document = _part(path, "word/document.xml")
    assert f'w:commentReference w:id="{reply["comment_id"]}"' not in document


# ---------------------------------------------------------------------------
# Documents that have no comments yet
# ---------------------------------------------------------------------------


def test_reading_a_document_with_no_comments_returns_empty(dc, tmp_path):
    path = _new_doc(tmp_path / "silent.docx")
    assert dc.list_comments(_reopen(path)) == []
    assert dc.list_threads(_reopen(path)) == []
    assert dc.thread_index(_reopen(path))["order"] == []


def test_threading_parts_are_created_on_a_document_that_never_had_comments(
        dc, validate, tmp_path):
    path = _new_doc(tmp_path / "fresh.docx")
    assert "word/comments.xml" not in _names(path)

    _commented(dc, path)
    names = _names(path)
    for part in ("word/comments.xml", "word/commentsExtended.xml",
                 "word/commentsIds.xml"):
        assert part in names, f"{part} was not created"

    content_types = _part(path, "[Content_Types].xml")
    for part in ("/word/commentsExtended.xml", "/word/commentsIds.xml"):
        assert f'PartName="{part}"' in content_types
    rels = _part(path, "word/_rels/document.xml.rels")
    assert "commentsExtended.xml" in rels
    assert "commentsIds.xml" in rels

    extended = _part(path, "word/commentsExtended.xml")
    assert "w15:paraId=" in extended
    assert 'w15:done="0"' in extended
    # The style ids the comment markup refers to must exist, or Word and
    # docx_validate.py both call the file broken.
    styles = _part(path, "word/styles.xml")
    assert 'w:styleId="CommentReference"' in styles
    assert 'w:styleId="CommentText"' in styles
    assert validate.validate(str(path))["ok"] is True


def test_paragraph_ids_are_unique_and_inside_word_s_range(dc, tmp_path):
    path = _new_doc(tmp_path / "ids.docx")
    root_id = _commented(dc, path)
    doc = _reopen(path)
    dc.reply_to_comment(doc, root_id, "One", "Hermes")
    dc.reply_to_comment(doc, root_id, "Two", "Hermes")
    doc.save(str(path))

    para_ids = [c["para_id"] for c in dc.list_comments(_reopen(path))]
    assert len(para_ids) == len(set(para_ids)) == 3
    for value in para_ids:
        assert len(value) == 8
        assert 0 < int(value, 16) < 0x80000000


def test_deleting_a_thread_root_leaves_its_replies_readable(dc, validate,
                                                            tmp_path):
    """A single-comment delete must not leave a reply pointing at nothing."""
    path = _new_doc(tmp_path / "root-gone.docx")
    root_id = _commented(dc, path)
    doc = _reopen(path)
    reply = dc.reply_to_comment(doc, root_id, "Added the source", "Hermes")
    doc.save(str(path))

    doc = _reopen(path)
    assert dc.delete_comment(doc, root_id) is True
    doc.save(str(path))

    threads = dc.list_threads(_reopen(path))
    assert [t["id"] for t in threads] == [reply["comment_id"]]
    assert threads[0]["parent_id"] is None
    assert "w15:paraIdParent" not in _part(path, "word/commentsExtended.xml")
    assert validate.validate(str(path))["ok"] is True


# ---------------------------------------------------------------------------
# Comments that predate threading
# ---------------------------------------------------------------------------


def _strip_threading(src: Path, dst: Path) -> Path:
    """Rewrite a document into the shape a pre-threading producer writes.

    No w14:paraId on any comment, no w14 declaration on the comments root,
    and none of the sibling parts. That is the file LibreOffice and older
    Word versions hand over, and it is the only shape that exercises the
    namespace rebuild and paraId minting.
    """
    import re

    drop = {"word/commentsExtended.xml", "word/commentsIds.xml"}
    with zipfile.ZipFile(str(src)) as zin, \
            zipfile.ZipFile(str(dst), "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            if item.filename in drop:
                continue
            data = zin.read(item.filename)
            if item.filename == "word/comments.xml":
                text = data.decode("utf-8")
                text = re.sub(r'\s+xmlns:w14="[^"]*"', "", text)
                text = re.sub(r'\s+w14:(paraId|textId)="[^"]*"', "", text)
                data = text.encode("utf-8")
            elif item.filename == "word/_rels/document.xml.rels":
                text = data.decode("utf-8")
                text = re.sub(
                    r'<Relationship[^>]*comments(Extended|Ids)\.xml"[^>]*/>',
                    "", text)
                data = text.encode("utf-8")
            elif item.filename == "[Content_Types].xml":
                text = data.decode("utf-8")
                text = re.sub(
                    r'<Override PartName="/word/comments(Extended|Ids)\.xml"'
                    r'[^>]*/>', "", text)
                data = text.encode("utf-8")
            zout.writestr(item, data)
    return dst


def test_a_comment_that_predates_threading_can_be_replied_to(dc, validate,
                                                             tmp_path):
    path = _new_doc(tmp_path / "modern.docx")
    root_id = _commented(dc, path)
    doc = _reopen(path)
    dc.reply_to_comment(doc, root_id, "Added the source", "Hermes")
    doc.save(str(path))
    legacy = _strip_threading(path, tmp_path / "legacy.docx")

    # Nothing to thread on yet, so every comment reads as its own root.
    before = dc.list_threads(_reopen(legacy))
    assert len(before) == 2
    assert all(t["para_id"] is None and t["replies"] == [] for t in before)

    doc = _reopen(legacy)
    reply = dc.reply_to_comment(doc, "0", "Second pass", "QA")
    doc.save(str(legacy))

    doc = _reopen(legacy)
    dc.set_thread_resolved(doc, "0", True)
    doc.save(str(legacy))

    comments = _part(legacy, "word/comments.xml")
    # A ns0-prefixed attribute next to mc:Ignorable="w14" is a dangling
    # prefix, which is exactly the file Word offers to repair.
    assert 'xmlns:w14="' in comments
    assert "ns0:" not in comments
    assert "w14:paraId=" in comments

    threads = {t["id"]: t for t in dc.list_threads(_reopen(legacy))}
    assert threads["0"]["resolved"] is True
    assert [r["id"] for r in threads["0"]["replies"]] == [
        reply["comment_id"]]
    assert threads["0"]["replies"][0]["resolved"] is True
    assert validate.validate(str(legacy))["ok"] is True


# ---------------------------------------------------------------------------
# Arabic
# ---------------------------------------------------------------------------


def test_arabic_comment_round_trips_with_its_text_intact(dc, validate,
                                                         tmp_path):
    body = "المراجعة النهائية للتقرير."
    path = _new_doc(tmp_path / "arabic.docx", body)
    comment = "يحتاج إلى مصدر"
    reply = "تم، أضفت المصدر"
    root_id = _commented(dc, path, target="المراجعة", text=comment,
                         author="سبار")

    doc = _reopen(path)
    dc.reply_to_comment(doc, root_id, reply, "هيرمس", "هـ")
    doc.save(str(path))

    threads = dc.list_threads(_reopen(path))
    assert threads[0]["text"] == comment
    assert threads[0]["author"] == "سبار"
    assert threads[0]["anchored_text"] == "المراجعة"
    assert threads[0]["replies"][0]["text"] == reply
    assert threads[0]["replies"][0]["author"] == "هيرمس"
    assert validate.validate(str(path))["ok"] is True


# ---------------------------------------------------------------------------
# The CLI
# ---------------------------------------------------------------------------


def _run(*args):
    return subprocess.run([sys.executable, str(SCRIPTS / "docx_comments.py"),
                           *args], capture_output=True, text=True,
                          cwd=str(SCRIPTS))


def test_cli_walks_a_whole_review_conversation(tmp_path):
    path = _new_doc(tmp_path / "cli.docx")
    added = _run("add", str(path), "--target", "Q3 revenue",
                 "--text", "Needs a source", "--author", "Reviewer")
    assert added.returncode == 0, added.stderr
    root_id = json.loads(added.stdout)["comment_id"]

    replied = _run("reply", str(path), "--id", root_id,
                   "--text", "Added the source", "--author", "Hermes")
    assert replied.returncode == 0, replied.stderr
    assert json.loads(replied.stdout)["parent_id"] == root_id

    resolved = _run("resolve", str(path), "--id", root_id)
    assert resolved.returncode == 0, resolved.stderr
    assert json.loads(resolved.stdout)["resolved"] is True

    listed = _run("list", str(path), "--json", "threads")
    assert listed.returncode == 0, listed.stderr
    threads = json.loads(listed.stdout)["threads"]
    assert len(threads) == 1 and threads[0]["resolved"] is True

    removed = _run("delete-thread", str(path), "--id", root_id)
    assert removed.returncode == 0, removed.stderr
    assert json.loads(_run("list", str(path)).stdout)["comments"] == []


def test_cli_reports_an_unknown_id_and_exits_non_zero(tmp_path):
    path = _new_doc(tmp_path / "missing.docx")
    for cmd in ("reply", "resolve", "reopen", "delete-thread"):
        args = [cmd, str(path), "--id", "99"]
        if cmd == "reply":
            args += ["--text", "hello"]
        result = _run(*args)
        assert result.returncode == 1, f"{cmd} should fail"
        payload = json.loads(result.stdout)
        assert payload["ok"] is False and "99" in payload["error"]


# ---------------------------------------------------------------------------
# LibreOffice round trip
# ---------------------------------------------------------------------------


@pytest.mark.skipif(shutil.which("soffice") is None,
                    reason="LibreOffice is not installed")
def test_threads_survive_a_libreoffice_round_trip(dc, validate, tmp_path):
    """LibreOffice keeps the thread and the resolved state.

    It drops word/commentsIds.xml and renumbers w:id, so the assertions key
    on author, text and resolved state rather than on the numeric id.
    """
    path = _new_doc(tmp_path / "trip.docx")
    root_id = _commented(dc, path)
    doc = _reopen(path)
    dc.reply_to_comment(doc, root_id, "Added the source", "Hermes")
    dc.reply_to_comment(doc, root_id, "تم التحقق", "Sbar")
    doc.save(str(path))
    doc = _reopen(path)
    dc.set_thread_resolved(doc, root_id, True)
    doc.save(str(path))

    outdir = tmp_path / "converted"
    profile = tmp_path / "loprofile"
    result = subprocess.run(
        ["soffice", "--headless",
         f"-env:UserInstallation=file://{profile}",
         "--convert-to", "docx", "--outdir", str(outdir), str(path)],
        capture_output=True, text=True, timeout=180)
    converted = outdir / "trip.docx"
    assert converted.exists(), result.stdout + result.stderr

    threads = dc.list_threads(_reopen(converted))
    assert len(threads) == 1
    thread = threads[0]
    assert thread["author"] == "Reviewer"
    assert thread["text"] == "Needs a source"
    assert thread["resolved"] is True
    assert [(r["author"], r["text"]) for r in thread["replies"]] == [
        ("Hermes", "Added the source"), ("Sbar", "تم التحقق")]
    assert all(r["resolved"] for r in thread["replies"])
    assert validate.validate(str(converted))["ok"] is True

    # The conversation is still writable after the round trip.
    doc = _reopen(converted)
    again = dc.reply_to_comment(doc, thread["id"], "Second pass", "QA")
    doc.save(str(converted))
    assert again["para_id_parent"] == thread["para_id"]
    assert len(dc.list_threads(_reopen(converted))[0]["replies"]) == 3
