"""Contracts for the house design system and its lint.

The system is applied by default by four create scripts, so the tests
that matter are the ones that keep it safe to apply by default: it must
never overwrite an explicit choice, it must leave the complex-script font
slot to the Arabic pass, and it must be reachable from each office skill
without an installed package.
"""

from __future__ import annotations

import importlib.util
import json
import subprocess
import sys
from pathlib import Path

import pytest

REPO = Path(__file__).resolve().parents[2]
HOUSE = REPO / "skills" / "productivity" / "house-style" / "scripts"
OFFICE = ["docx", "xlsx", "pdf", "powerpoint"]


def _load(name: str, path: Path):
    spec = importlib.util.spec_from_file_location(name, path)
    module = importlib.util.module_from_spec(spec)
    sys.modules[name] = module
    spec.loader.exec_module(module)
    return module


@pytest.fixture(scope="module")
def hs():
    if str(HOUSE) not in sys.path:
        sys.path.insert(0, str(HOUSE))
    return _load("house_style_under_test", HOUSE / "house_style.py")


@pytest.fixture(scope="module")
def lint():
    if str(HOUSE) not in sys.path:
        sys.path.insert(0, str(HOUSE))
    return _load("style_lint_under_test", HOUSE / "style_lint.py")


# ---------------------------------------------------------------------------
# The system itself
# ---------------------------------------------------------------------------


@pytest.mark.parametrize("name", ["editorial", "slate", "mono"])
def test_every_preset_clears_wcag_aa(hs, name):
    """Body text, captions and accent text all read on paper and surface."""
    theme = hs.load_theme(name)
    p = theme.palette
    for fg, bg in ((p.ink, p.paper), (p.muted, p.paper), (p.ink, p.surface),
                   (p.accent_text, p.paper)):
        assert hs.contrast_ratio(fg, bg) >= 4.5, f"{name}: {fg} on {bg}"


def test_deck_title_is_twice_its_body(hs):
    theme = hs.load_theme()
    assert theme.deck["title"] / theme.deck["body"] >= 2.0
    assert theme.deck["body"] >= 18       # the projected floor
    assert theme.doc["body"] == 11


def test_accent_retint_keeps_neutrals_and_rebuilds_the_rest(hs):
    base = hs.load_theme("slate")
    tinted = hs.load_theme("slate", accent="B4482E")
    assert tinted.palette.accent == "B4482E"
    assert tinted.palette.ink == base.palette.ink
    assert tinted.palette.paper == base.palette.paper
    assert tinted.palette.series[0] == "B4482E"
    # The accent as chosen is usually too light to set text in.
    assert hs.contrast_ratio(tinted.palette.accent_text,
                             tinted.palette.paper) >= 4.5


def test_theme_from_spec_reads_every_shape(hs):
    assert hs.theme_from_spec({"theme": False}) is None
    assert hs.theme_from_spec({"theme": "mono"}).name == "mono"
    assert hs.theme_from_spec({}).name == "editorial"
    custom = hs.theme_from_spec({"theme": {"name": "slate", "accent": "112233"}})
    assert custom.palette.accent == "112233"
    with pytest.raises(ValueError):
        hs.theme_from_spec({"theme": 7})


def test_arabic_detection_and_leading(hs):
    assert hs.is_arabic("نمت الإيرادات")
    assert not hs.is_arabic("revenue grew 34 percent")
    assert not hs.is_arabic("2026-09-11")
    assert hs.line_spacing_for("نص عربي") == 1.45
    assert hs.line_spacing_for("latin text", latin=1.25) == 1.25
    # A title wants the opposite of a paragraph.
    assert hs.line_spacing_for("عنوان", arabic=1.25) == 1.25


# ---------------------------------------------------------------------------
# The appliers
# ---------------------------------------------------------------------------


def test_docx_grid_rules_every_column_and_row(hs):
    """The default table is a full grid in a quiet grey."""
    pytest.importorskip("docx")
    from docx import Document
    from docx.oxml.ns import qn

    doc = Document()
    table = doc.add_table(rows=3, cols=3)
    table.cell(0, 0).text = "Product"
    theme = hs.load_theme()
    hs.theme_docx(doc, theme)

    borders = table._tbl.tblPr.find(qn("w:tblBorders"))
    for edge in ("insideV", "insideH", "left", "right", "top", "bottom"):
        node = borders.find(qn(f"w:{edge}"))
        assert node is not None and node.get(qn("w:val")) == "single", edge
        assert node.get(qn("w:color")) == theme.palette.grid, edge


def test_docx_rules_style_is_still_booktabs(hs):
    pytest.importorskip("docx")
    from docx import Document
    from docx.oxml.ns import qn

    doc = Document()
    doc.add_table(rows=2, cols=2)
    hs.theme_docx(doc, hs.load_theme(table_style="rules"))

    borders = doc.tables[0]._tbl.tblPr.find(qn("w:tblBorders"))
    for edge in ("insideV", "left", "right", "insideH"):
        node = borders.find(qn(f"w:{edge}"))
        assert node is not None and node.get(qn("w:val")) == "none", edge


def test_docx_pass_respects_an_explicit_size_and_quiets_the_table(hs):
    pytest.importorskip("docx")
    from docx import Document
    from docx.oxml.ns import qn
    from docx.shared import Pt

    doc = Document()
    run = doc.add_paragraph().add_run("deliberately large")
    run.font.size = Pt(37)
    table = doc.add_table(rows=2, cols=2)
    table.style = "Table Grid"
    table.cell(0, 0).text = "Product"

    hs.theme_docx(doc, hs.load_theme(table_style="rules"))

    assert run.font.size == Pt(37), "an explicit size must survive the pass"
    assert doc.styles["Normal"].font.size == Pt(11)
    borders = table._tbl.tblPr.find(qn("w:tblBorders"))
    assert borders is not None
    for edge in ("insideV", "left", "right"):
        node = borders.find(qn(f"w:{edge}"))
        assert node is not None and node.get(qn("w:val")) == "none"


def test_docx_pass_sets_arabic_leading_and_drops_italics(hs):
    pytest.importorskip("docx")
    from docx import Document

    doc = Document()
    para = doc.add_paragraph()
    run = para.add_run("هذه فقرة عربية كاملة للتجربة")
    run.font.italic = True

    hs.theme_docx(doc, hs.load_theme())

    assert para.paragraph_format.line_spacing == 1.45
    assert run.font.italic is False, "Arabic carries emphasis in weight"
    assert run.font.bold is True


def test_docx_pass_leaves_the_complex_script_slot_to_the_arabic_pass(hs):
    """The two passes must not fight over w:cs. This one owns w:ascii."""
    pytest.importorskip("docx")
    from docx import Document
    from docx.oxml.ns import qn

    doc = Document()
    doc.add_paragraph("نص عربي")
    hs.theme_docx(doc, hs.load_theme())

    rpr = doc.styles["Normal"].element.find(qn("w:rPr"))
    fonts = rpr.find(qn("w:rFonts")) if rpr is not None else None
    if fonts is not None:
        assert fonts.get(qn("w:cs")) is None


def test_pptx_pass_sizes_the_title_and_keeps_the_body_off_a_chart(hs):
    pytest.importorskip("pptx")
    from pptx import Presentation
    from pptx.util import Inches, Pt

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    cover = prs.slides.add_slide(prs.slide_layouts[0])
    cover.shapes.title.text = "Q3 review"
    cover.placeholders[1].text = "Sales and margin"

    body_slide = prs.slides.add_slide(prs.slide_layouts[1])
    body_slide.shapes.title.text = "Revenue grew"
    body_slide.placeholders[1].text_frame.text = "Handles carried the quarter"
    body_slide.shapes.add_textbox(Inches(7.2), Inches(1.9), Inches(5),
                                  Inches(3)).text_frame.text = "chart stand-in"

    theme = hs.load_theme()
    hs.theme_pptx(prs, theme)

    cover_run = cover.shapes.title.text_frame.paragraphs[0].runs[0]
    title_run = body_slide.shapes.title.text_frame.paragraphs[0].runs[0]
    assert cover_run.font.size == Pt(theme.deck["cover_title"])
    assert title_run.font.size == Pt(theme.deck["title"])

    body = body_slide.placeholders[1]
    right_edge = (body.left + body.width) / 914400
    assert right_edge <= 7.2 - theme.geometry["gutter"] + 0.01, \
        "the body column must stop before the shape the spec placed"


def test_a_long_title_steps_down_the_scale_and_the_body_clears_it(hs):
    """The overflow guard, in both scripts.

    fit_title is the documented reason a long title does not print over
    the first bullet, and until this test existed nothing called it with
    a title long enough to step down, so the whole ladder could have been
    deleted without a single failure.
    """
    pytest.importorskip("pptx")
    from pptx import Presentation
    from pptx.util import Inches, Pt

    theme = hs.load_theme()
    long_title = ("Revenue grew thirty four percent across two product "
                  "families while gaskets fell after the July price move")
    size, lines = hs.fit_title(long_title, 12.0, theme.deck)
    assert size < theme.deck["title"], "a twenty word title must step down"
    assert lines <= 2
    short = hs.fit_title("Revenue grew", 12.0, theme.deck)
    assert short[0] == theme.deck["title"], "a short title keeps its size"

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = long_title
    slide.placeholders[1].text_frame.text = "Handles carried the quarter"
    hs.theme_pptx(prs, theme)

    title = slide.shapes.title
    body = slide.placeholders[1]
    assert title.text_frame.paragraphs[0].runs[0].font.size == Pt(size)
    title_bottom = (title.top + title.height) / 914400
    assert body.top / 914400 >= title_bottom - 0.01, \
        "the body must start below the title box, not inside it"


def test_an_arabic_long_title_gets_its_own_leading_and_still_clears(hs):
    """Arabic leading is looser, so the same title needs a taller band."""
    pytest.importorskip("pptx")
    from pptx import Presentation
    from pptx.util import Inches

    theme = hs.load_theme()
    arabic = ("نمت الإيرادات أربعة وثلاثين بالمئة على عائلتين من المنتجات "
              "بينما تراجعت الحشوات بعد تعديل السعر في تموز")
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = arabic
    slide.placeholders[1].text_frame.text = "المقابض حملت الربع"
    hs.theme_pptx(prs, theme)

    title, body = slide.shapes.title, slide.placeholders[1]
    assert (body.top / 914400) >= (title.top + title.height) / 914400 - 0.01
    assert title.text_frame.paragraphs[0].line_spacing == 1.25, \
        "an Arabic title takes the tighter display leading, not the body one"


def test_a_sized_run_inside_a_table_cell_is_left_alone(hs):
    """The one place the docx pass writes a run size directly.

    The old test asserted this on a body paragraph, where the pass never
    touches runs at all, so the assertion could not fail and the real
    path went unguarded.
    """
    pytest.importorskip("docx")
    from docx import Document
    from docx.shared import Pt, RGBColor

    doc = Document()
    table = doc.add_table(rows=2, cols=2)
    cell = table.cell(1, 0)
    run = cell.paragraphs[0].add_run("deliberately large")
    run.font.size = Pt(19)
    run.font.color.rgb = RGBColor.from_string("FF00FF")
    run.font.bold = False
    header = table.cell(0, 0).paragraphs[0].add_run("Product")

    hs.theme_docx(doc, hs.load_theme())

    assert run.font.size == Pt(19), "a sized cell run must survive the pass"
    assert str(run.font.color.rgb) == "FF00FF"
    assert run.font.bold is False, "an explicit weight is a decision too"
    assert header.font.size is not None, "the header still gets the house size"


def test_pptx_pass_respects_explicit_run_formatting(hs):
    pytest.importorskip("pptx")
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.util import Pt

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Title"
    frame = slide.placeholders[1].text_frame
    frame.text = "deliberate"
    run = frame.paragraphs[0].runs[0]
    run.font.size = Pt(31)
    run.font.color.rgb = RGBColor.from_string("FF00FF")

    hs.theme_pptx(prs, hs.load_theme())

    assert run.font.size == Pt(31)
    assert str(run.font.color.rgb) == "FF00FF"


def test_pptx_pass_places_an_arabic_body_against_the_right_margin(hs):
    pytest.importorskip("pptx")
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "نمت الإيرادات"
    slide.placeholders[1].text_frame.text = "المقابض حملت الربع"

    theme = hs.load_theme()
    hs.theme_pptx(prs, theme)

    body = slide.placeholders[1]
    right_gap = 13.333 - (body.left + body.width) / 914400
    assert abs(right_gap - theme.geometry["margin_x"]) < 0.02


def test_xlsx_pass_quiets_the_sheet(hs):
    pytest.importorskip("openpyxl")
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.append(["Region", "Revenue"])
    for i in range(5):
        ws.append([f"r{i}", i])

    hs.theme_xlsx(wb, hs.load_theme())

    assert ws.sheet_view.showGridLines is False, "Excel's own grid stays off"
    assert ws.freeze_panes == "A2"
    assert ws["A1"].font.bold is True
    assert ws["A1"].border.bottom.style == "thin"
    # The data region carries its own ruling, in the house grey.
    assert ws["A2"].border.left.style == "thin"
    assert ws["A2"].border.left.color.rgb.endswith(hs.load_theme().palette.grid)

    plain = Workbook()
    sheet = plain.active
    sheet.append(["Region", "Revenue"])
    sheet.append(["North", 100])
    hs.theme_xlsx(plain, hs.load_theme(table_style="rules"))
    assert sheet["A2"].border.left.style is None, "rules style stays unruled"


def test_xlsx_pass_respects_an_explicit_color_and_size(hs):
    pytest.importorskip("openpyxl")
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.active
    ws.append(["Region", "Revenue"])
    ws.append(["North", 100])
    ws["A1"].font = Font(color="FF00FF00", size=18)
    ws["A2"].font = Font(color="FFFF0000")

    hs.theme_xlsx(wb, hs.load_theme())

    assert ws["A1"].font.color.rgb == "FF00FF00", "a header color must survive"
    assert ws["A1"].font.size == 18
    assert ws["A2"].font.color.rgb == "FFFF0000"


def test_pptx_grid_rules_the_table_cells(hs):
    pytest.importorskip("pptx")
    from pptx import Presentation
    from pptx.oxml.ns import qn
    from pptx.util import Inches

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    table = slide.shapes.add_table(3, 2, Inches(1), Inches(1), Inches(6),
                                   Inches(2)).table
    table.cell(0, 0).text = "Product"
    theme = hs.load_theme()
    hs.theme_pptx(prs, theme)

    tc_pr = table.rows[1].cells[0]._tc.get_or_add_tcPr()
    for tag in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
        line = tc_pr.find(qn(tag))
        assert line is not None, tag
        color = line.find(qn("a:solidFill")).find(qn("a:srgbClr"))
        assert color.get("val") == theme.palette.grid


def test_xlsx_pass_keeps_a_border_the_spec_drew(hs):
    pytest.importorskip("openpyxl")
    from openpyxl import Workbook
    from openpyxl.styles import Border, Side

    wb = Workbook()
    ws = wb.active
    ws.append(["Region", "Revenue"])
    ws.append(["North", 100])
    ws["A1"].border = Border(left=Side(style="medium"),
                             bottom=Side(style="medium"))

    hs.theme_xlsx(wb, hs.load_theme())

    assert ws["A1"].border.left.style == "medium", "a drawn border is a decision"
    assert ws["A1"].border.bottom.style == "medium"


def test_xlsx_pass_flips_an_arabic_sheet(hs):
    pytest.importorskip("openpyxl")
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.append(["المنتج", "الهامش"])
    ws.append(["المقابض", "الحشوات"])

    hs.theme_xlsx(wb, hs.load_theme())
    assert ws.sheet_view.rightToLeft is True


def test_pdf_styles_replace_the_reportlab_sample(hs):
    pytest.importorskip("reportlab")
    styles, table_style = hs.pdf_styles(hs.load_theme())
    assert styles["Heading1"].fontSize == hs.load_theme().doc["h1"]
    commands = table_style(4).getCommands()
    assert not any(cmd[0] == "GRID" for cmd in commands), "no boxed grid"
    assert any(cmd[0] == "LINEBELOW" for cmd in commands)


# ---------------------------------------------------------------------------
# The lint
# ---------------------------------------------------------------------------


def test_lint_passes_a_house_grid_and_flags_a_black_one(hs, lint, tmp_path):
    """A ruled table is the house default. A near black grid is not."""
    pytest.importorskip("docx")
    from docx import Document

    housed = tmp_path / "housed.docx"
    doc = Document()
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Product"
    table.cell(1, 0).text = "Handles"
    hs.theme_docx(doc, hs.load_theme())
    doc.save(str(housed))
    assert "loud_table_rules" not in {f["rule"] for f in lint.lint(housed)}

    stock = tmp_path / "stock.docx"
    plain = Document()
    plain.add_table(rows=2, cols=2).style = "Table Grid"
    plain.save(str(stock))
    assert "loud_table_rules" in {f["rule"] for f in lint.lint(stock)}


def test_lint_fails_on_a_long_dash(lint, tmp_path):
    draft = tmp_path / "draft.md"
    draft.write_text("A sentence \u2014 with an em dash.\n", encoding="utf-8")
    findings = lint.lint(draft)
    dash = [f for f in findings if f["rule"] == "long_dash"]
    assert dash and dash[0]["level"] == "error"


def test_lint_flags_machine_prose_as_warnings_only(lint, tmp_path):
    draft = tmp_path / "draft.md"
    draft.write_text(
        "In today's fast-paced world, the team aims to delve into results.\n",
        encoding="utf-8")
    findings = lint.lint(draft)
    rules = {f["rule"] for f in findings}
    assert {"llm_lexicon", "intention_not_action", "stock_opener"} <= rules
    assert all(f["level"] == "warn" for f in findings)


def test_lint_passes_clean_arabic(lint, tmp_path):
    draft = tmp_path / "clean.md"
    draft.write_text("نمت الإيرادات 34 بالمئة. المقابض حملت الربع.\n",
                     encoding="utf-8")
    assert lint.lint(draft) == []


def test_lint_reads_a_themed_deck_without_complaint(hs, lint, tmp_path):
    pytest.importorskip("pptx")
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Revenue grew 34 percent"
    slide.placeholders[1].text_frame.text = "Handles carried the quarter"
    slide.shapes.add_table(2, 2, Inches(1), Inches(4), Inches(5), Inches(1))
    hs.theme_pptx(prs, hs.load_theme())
    out = tmp_path / "deck.pptx"
    prs.save(str(out))

    rules = {f["rule"] for f in lint.lint(out)}
    assert "stock_font" not in rules
    assert "type_too_small" not in rules


# ---------------------------------------------------------------------------
# Wiring: the office skills reach the system, and say so in their output
# ---------------------------------------------------------------------------


@pytest.mark.parametrize("name", OFFICE)
def test_every_office_skill_carries_the_locator(name):
    bridge = REPO / "skills" / "productivity" / name / "scripts" / "house_common.py"
    assert bridge.exists(), f"{name}: no house_common.py"
    assert bridge.read_text(encoding="utf-8") == \
        (HOUSE / "house_common.py").read_text(encoding="utf-8"), \
        f"{name}: the locator has drifted from the house-style copy"


def test_locator_returns_none_when_the_system_is_absent(tmp_path):
    """An office skill installed alone still builds files."""
    lone = tmp_path / "scripts"
    lone.mkdir()
    src = (HOUSE / "house_common.py").read_text(encoding="utf-8")
    (lone / "house_common.py").write_text(src, encoding="utf-8")
    probe = (
        "import sys, json;"
        f"sys.path.insert(0, {str(lone)!r});"
        "from house_common import load_house_style;"
        "print(json.dumps(load_house_style() is None))"
    )
    out = subprocess.run([sys.executable, "-c", probe], capture_output=True,
                         text=True, env={"HOME": str(tmp_path), "PATH": ""})
    assert out.stdout.strip() == "true", out.stderr


def test_deck_creation_reports_the_theme_it_applied(tmp_path):
    pytest.importorskip("pptx")
    spec = tmp_path / "deck.json"
    spec.write_text(json.dumps({"slides": [
        {"layout": "title", "title": "Q3", "subtitle": "review"}]}),
        encoding="utf-8")
    out = tmp_path / "deck.pptx"
    script = REPO / "skills" / "productivity" / "powerpoint" / "scripts" / \
        "pptx_create.py"
    run = subprocess.run([sys.executable, str(script), str(spec), str(out)],
                         capture_output=True, text=True)
    assert run.returncode == 0, run.stderr
    assert json.loads(run.stdout)["theme"] == "editorial"

    run_off = subprocess.run(
        [sys.executable, str(script), str(spec), str(out), "--no-theme"],
        capture_output=True, text=True)
    assert json.loads(run_off.stdout)["theme"] is None
