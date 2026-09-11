"""Contracts for comments in decks and workbooks.

Both formats keep their comments in parts the libraries do not model, so
these scripts write the XML themselves. That makes two things worth
testing above everything else: a file that still opens, and a thread that
still reads as a thread after a round trip through the library that did
not write it.
"""

from __future__ import annotations

import json
import subprocess
import sys
import zipfile
from pathlib import Path

import pytest

REPO = Path(__file__).resolve().parents[2]
PPTX_CLI = REPO / "skills" / "productivity" / "powerpoint" / "scripts" / \
    "pptx_comments.py"
XLSX_CLI = REPO / "skills" / "productivity" / "xlsx" / "scripts" / \
    "xlsx_comments.py"


def run(cli: Path, *args: str) -> dict:
    out = subprocess.run([sys.executable, str(cli), *args],
                         capture_output=True, text=True)
    assert out.returncode == 0, f"{args}: {out.stderr}"
    return json.loads(out.stdout)


def fails(cli: Path, *args: str) -> subprocess.CompletedProcess:
    return subprocess.run([sys.executable, str(cli), *args],
                          capture_output=True, text=True)


@pytest.fixture
def deck(tmp_path):
    pytest.importorskip("pptx")
    from pptx import Presentation

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Revenue"
    slide.placeholders[1].text_frame.text = "Handles carried the quarter"
    out = tmp_path / "deck.pptx"
    prs.save(str(out))
    return out


@pytest.fixture
def book(tmp_path):
    pytest.importorskip("openpyxl")
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.append(["Region", "Revenue"])
    ws.append(["North", 100])
    out = tmp_path / "book.xlsx"
    wb.save(str(out))
    return out


# ---------------------------------------------------------------------------
# Decks
# ---------------------------------------------------------------------------


def test_deck_thread_reads_back_with_its_reply_and_state(deck):
    assert run(PPTX_CLI, "list", str(deck))["comments"] == []

    added = run(PPTX_CLI, "add", str(deck), "--slide", "1",
                "--text", "Is this the closing figure?", "--author", "Reviewer")
    parent = added["id"]
    run(PPTX_CLI, "reply", str(deck), "--id", parent,
        "--text", "Yes, from the Q3 close.", "--author", "Hermes")

    comments = run(PPTX_CLI, "list", str(deck))["comments"]
    assert len(comments) == 2
    reply = next(c for c in comments if c["parent_id"])
    assert reply["parent_id"] == parent
    assert reply["author"] == "Hermes"
    assert all(c["resolved"] is False for c in comments)

    run(PPTX_CLI, "resolve", str(deck), "--id", parent)
    assert all(c["resolved"] for c in run(PPTX_CLI, "list",
                                          str(deck))["comments"]), \
        "resolving a thread resolves its replies too"

    run(PPTX_CLI, "reopen", str(deck), "--id", parent)
    assert not any(c["resolved"] for c in run(PPTX_CLI, "list",
                                              str(deck))["comments"])


def test_deck_delete_takes_the_replies_with_it(deck):
    parent = run(PPTX_CLI, "add", str(deck), "--slide", "1",
                 "--text", "Cut this slide?", "--author", "Reviewer")["id"]
    run(PPTX_CLI, "reply", str(deck), "--id", parent, "--text", "Agreed.",
        "--author", "Hermes")
    run(PPTX_CLI, "delete", str(deck), "--id", parent)
    assert run(PPTX_CLI, "list", str(deck))["comments"] == []


def test_deck_arabic_comment_survives(deck):
    run(PPTX_CLI, "add", str(deck), "--slide", "1",
        "--text", "هل الرقم محدث؟", "--author", "محمد")
    comment = run(PPTX_CLI, "list", str(deck))["comments"][0]
    assert comment["text"] == "هل الرقم محدث؟"
    assert comment["author"] == "محمد"


def test_deck_stays_a_valid_package_and_still_opens(deck):
    pytest.importorskip("pptx")
    from pptx import Presentation

    run(PPTX_CLI, "add", str(deck), "--slide", "1", "--text", "note",
        "--author", "Reviewer")
    with zipfile.ZipFile(deck) as zf:
        assert zf.testzip() is None
        names = set(zf.namelist())
        types = zf.read("[Content_Types].xml").decode("utf-8")
        for part in names:
            if "comments/" in part or part.endswith("authors.xml"):
                assert f"/{part}" in types, f"{part} has no content type"
    # The library that did not write the part must still open the file.
    assert len(Presentation(str(deck)).slides._sldIdLst) == 1


def test_deck_unknown_thread_is_an_error_not_a_silent_pass(deck):
    out = fails(PPTX_CLI, "reply", str(deck), "--id", "{no-such-id}",
                "--text", "hello", "--author", "Hermes")
    assert out.returncode != 0


# ---------------------------------------------------------------------------
# Workbooks
# ---------------------------------------------------------------------------


def test_workbook_thread_reads_back_with_its_reply_and_state(book):
    assert run(XLSX_CLI, "list", str(book))["comments"] == []

    parent = run(XLSX_CLI, "add", str(book), "--sheet", "Sheet",
                 "--cell", "B2", "--text", "Stale figure?",
                 "--author", "Reviewer")["id"]
    run(XLSX_CLI, "reply", str(book), "--id", parent,
        "--text", "Refreshed today.", "--author", "Hermes")

    comments = run(XLSX_CLI, "list", str(book))["comments"]
    assert len(comments) == 2
    assert {c["cell"] for c in comments} == {"B2"}
    assert any(c["parent_id"] == parent for c in comments)

    run(XLSX_CLI, "resolve", str(book), "--id", parent)
    assert all(c["resolved"] for c in run(XLSX_CLI, "list",
                                          str(book))["comments"])
    run(XLSX_CLI, "reopen", str(book), "--id", parent)
    assert not any(c["resolved"] for c in run(XLSX_CLI, "list",
                                              str(book))["comments"])


def test_workbook_arabic_comment_survives(book):
    run(XLSX_CLI, "add", str(book), "--sheet", "Sheet", "--cell", "A2",
        "--text", "الرقم قديم", "--author", "محمد")
    comment = run(XLSX_CLI, "list", str(book))["comments"][0]
    assert comment["text"] == "الرقم قديم"


def test_workbook_keeps_its_chart_through_a_comment_edit(tmp_path):
    """The parts are written into the zip by hand, so anything openpyxl
    does not model is exactly what a careless implementation drops."""
    pytest.importorskip("openpyxl")
    from openpyxl import Workbook, load_workbook
    from openpyxl.chart import BarChart, Reference

    wb = Workbook()
    ws = wb.active
    ws.append(["Region", "Revenue"])
    ws.append(["North", 100])
    ws.append(["South", 80])
    chart = BarChart()
    chart.add_data(Reference(ws, min_col=2, min_row=1, max_row=3),
                   titles_from_data=True)
    ws.add_chart(chart, "D2")
    out = tmp_path / "chart.xlsx"
    wb.save(str(out))

    run(XLSX_CLI, "add", str(out), "--sheet", "Sheet", "--cell", "A2",
        "--text", "check this", "--author", "Reviewer")

    reopened = load_workbook(str(out))
    assert len(reopened.worksheets[0]._charts) == 1
    assert reopened.worksheets[0]["B2"].value == 100


def test_workbook_legacy_note_still_works_beside_the_threads(book):
    run(XLSX_CLI, "add-note", str(book), "--sheet", "Sheet", "--cell", "A1",
        "--text", "legacy note", "--author", "Reviewer")
    run(XLSX_CLI, "add", str(book), "--sheet", "Sheet", "--cell", "B2",
        "--text", "threaded", "--author", "Reviewer")
    kinds = {c["kind"] for c in run(XLSX_CLI, "list", str(book))["comments"]}
    assert "note" in kinds and "threaded" in kinds


def test_workbook_stays_a_valid_package(book):
    pytest.importorskip("openpyxl")
    from openpyxl import load_workbook

    run(XLSX_CLI, "add", str(book), "--sheet", "Sheet", "--cell", "B2",
        "--text", "note", "--author", "Reviewer")
    with zipfile.ZipFile(book) as zf:
        assert zf.testzip() is None
        types = zf.read("[Content_Types].xml").decode("utf-8")
        for part in zf.namelist():
            if "threadedComment" in part or part.endswith("person.xml"):
                assert f"/{part}" in types, f"{part} has no content type"
    assert load_workbook(str(book)).worksheets[0]["A1"].value == "Region"
