"""Contracts for embedding fonts into a deck.

A font name is not a font. The deck carries a family name and the
reader's machine resolves it, which for Arabic means tofu on any machine
without the face installed. pptx_embed_fonts.py puts the font bytes in
the package, in parts python-pptx does not model, so the two things
worth testing above everything else are a package that is still
internally consistent and a deck that still opens.

The third is the licence gate. Embedding ships the font inside every
copy of the deliverable, so a face with no licence file on disk is
refused unless the caller says otherwise, and that refusal has to keep
working.

Most tests build their own font directory out of one .ttf copied under a
synthetic family name, with and without an OFL.txt beside it. That keeps
the gate deterministic instead of depending on which faces this
particular machine happens to have and where its distribution files
their licences.
"""

from __future__ import annotations

import json
import os
import shutil
import subprocess
import sys
import zipfile
from pathlib import Path

import pytest

REPO = Path(__file__).resolve().parents[2]
CLI = REPO / "skills" / "productivity" / "powerpoint" / "scripts" / \
    "pptx_embed_fonts.py"

A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
P_NS = "http://schemas.openxmlformats.org/presentationml/2006/main"
R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
CT_NS = "http://schemas.openxmlformats.org/package/2006/content-types"
REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
CT_FONTDATA = "application/x-fontdata"

FAMILY = "HermesTestFace"
ARABIC_FAMILY = "IBM Plex Sans Arabic"
ARABIC_TEXT = "الإيرادات ارتفعت بنسبة 12 بالمئة"


def run(*args: str, fonts: Path | None = None) -> dict:
    out = _invoke(args, fonts)
    assert out.returncode == 0, f"{args}: {out.stderr}"
    return json.loads(out.stdout)


def _invoke(args, fonts):
    env = dict(os.environ)
    if fonts is not None:
        env["PPTX_FONT_DIRS"] = str(fonts)
    else:
        env.pop("PPTX_FONT_DIRS", None)
    return subprocess.run([sys.executable, str(CLI), *args],
                          capture_output=True, text=True, env=env)


# ---------------------------------------------------------------------------
# Fixtures
# ---------------------------------------------------------------------------


def _some_ttf() -> Path | None:
    """Any TrueType face on this machine, to copy under a test name.

    The bytes only have to be a real font; nothing in the embedding path
    reads the family out of the file, so a copy renamed to the synthetic
    family exercises exactly the code a real face would.
    """
    try:
        out = subprocess.run(["fc-match", "-f", "%{file}", "sans"],
                             capture_output=True, text=True, timeout=10)
    except (OSError, subprocess.SubprocessError):
        return None
    path = Path(out.stdout.strip()) if out.returncode == 0 else None
    if path is not None and path.suffix.lower() == ".ttf" and path.is_file():
        return path
    for base in (Path("/usr/share/fonts"), Path.home() / ".local/share/fonts"):
        if base.is_dir():
            for hit in base.rglob("*.ttf"):
                return hit
    return None


@pytest.fixture
def ttf():
    path = _some_ttf()
    if path is None:
        pytest.skip("no .ttf on this machine to build a test family from")
    return path


@pytest.fixture
def licensed_fonts(tmp_path, ttf):
    """A font directory holding the test family and its OFL licence."""
    root = tmp_path / "licensed"
    (root / FAMILY).mkdir(parents=True)
    shutil.copy(ttf, root / FAMILY / f"{FAMILY}-Regular.ttf")
    shutil.copy(ttf, root / FAMILY / f"{FAMILY}-Bold.ttf")
    (root / FAMILY / "OFL.txt").write_text(
        "SIL OPEN FONT LICENSE Version 1.1\n", encoding="utf-8")
    return root


@pytest.fixture
def unlicensed_fonts(tmp_path, ttf):
    """The same family with no licence file anywhere above it."""
    root = tmp_path / "unlicensed"
    (root / FAMILY).mkdir(parents=True)
    shutil.copy(ttf, root / FAMILY / f"{FAMILY}-Regular.ttf")
    return root


def _deck(tmp_path, name, family, text):
    pytest.importorskip("pptx")
    from lxml import etree
    from pptx import Presentation

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = text
    slide.placeholders[1].text_frame.text = text
    if family:
        for shape in (slide.shapes.title, slide.placeholders[1]):
            for para in shape.text_frame.paragraphs:
                for run_ in para.runs:
                    r_pr = run_._r.get_or_add_rPr()
                    for tag in ("latin", "cs"):
                        element = r_pr.find(f"{{{A_NS}}}{tag}")
                        if element is None:
                            element = etree.SubElement(
                                r_pr, f"{{{A_NS}}}{tag}")
                        element.set("typeface", family)
    out = tmp_path / name
    prs.save(str(out))
    return out


@pytest.fixture
def deck(tmp_path):
    """A deck whose runs name the synthetic test family."""
    return _deck(tmp_path, "deck.pptx", FAMILY, "Quarter three")


@pytest.fixture
def latin_deck(tmp_path):
    """A deck that names no family of its own, only the theme's."""
    return _deck(tmp_path, "latin.pptx", None, "Handles carried the quarter")


# ---------------------------------------------------------------------------
# Helpers that read the package the way a reader would
# ---------------------------------------------------------------------------


def package(path: Path) -> dict:
    from lxml import etree

    with zipfile.ZipFile(path) as z:
        names = set(z.namelist())
        presentation = etree.fromstring(z.read("ppt/presentation.xml"))
        rels = etree.fromstring(z.read("ppt/_rels/presentation.xml.rels"))
        types = etree.fromstring(z.read("[Content_Types].xml"))
    overrides = {ov.get("PartName"): ov.get("ContentType")
                 for ov in types.findall(f"{{{CT_NS}}}Override")}
    targets = {}
    for rel in rels.findall(f"{{{REL_NS}}}Relationship"):
        targets[rel.get("Id")] = "ppt/" + rel.get("Target").lstrip("/")
    return {"names": names, "presentation": presentation,
            "overrides": overrides, "targets": targets}


def embedded_entries(presentation) -> dict:
    """Family to slot to relationship id, straight out of the XML."""
    out = {}
    lst = presentation.find(f"{{{P_NS}}}embeddedFontLst")
    if lst is None:
        return out
    for entry in lst.findall(f"{{{P_NS}}}embeddedFont"):
        font = entry.find(f"{{{P_NS}}}font")
        slots = {}
        for slot in ("regular", "bold", "italic", "boldItalic"):
            element = entry.find(f"{{{P_NS}}}{slot}")
            if element is not None:
                slots[slot] = element.get(f"{{{R_NS}}}id")
        out[font.get("typeface")] = slots
    return out


def assert_consistent(path: Path):
    """Every reference resolves, every part is typed, nothing orphaned."""
    pkg = package(path)
    referenced = set()
    for family, slots in embedded_entries(pkg["presentation"]).items():
        assert slots, f"{family} has no face parts"
        for slot, rel_id in slots.items():
            target = pkg["targets"].get(rel_id)
            assert target, f"{family} {slot}: {rel_id} has no relationship"
            assert target in pkg["names"], f"{target} is not in the package"
            assert pkg["overrides"].get("/" + target) == CT_FONTDATA
            referenced.add(target)
    for name in pkg["names"]:
        if name.startswith("ppt/fonts/") and name.endswith(".fntdata"):
            assert name in referenced, f"{name} is orphaned"

    # The schema fixes where p:embeddedFontLst may sit, and out of place
    # is exactly what makes PowerPoint offer to repair the file.
    order = [child.tag.split("}")[-1] for child in pkg["presentation"]]
    if "embeddedFontLst" in order:
        at = order.index("embeddedFontLst")
        for early in ("sldMasterIdLst", "sldSz", "notesSz"):
            if early in order:
                assert order.index(early) < at
        if "defaultTextStyle" in order:
            assert at < order.index("defaultTextStyle")


def reopens(path: Path):
    from pptx import Presentation

    prs = Presentation(str(path))
    return [slide.shapes.title.text for slide in prs.slides]


# ---------------------------------------------------------------------------
# What the deck uses
# ---------------------------------------------------------------------------


def test_report_names_the_family_the_runs_ask_for(deck, licensed_fonts):
    report = run("report", str(deck), fonts=licensed_fonts)
    assert FAMILY in report["families_used"]
    assert "slides" in report["families_used"][FAMILY]
    assert report["dry_run"] is True
    # A dry run writes nothing.
    assert report["bytes_added"] > 0
    assert "ppt/fonts/font1.fntdata" not in package(deck)["names"]


def test_report_resolves_the_theme_tokens_rather_than_embedding_them(
        latin_deck, licensed_fonts):
    """A slide almost never names a family: it points at the theme.

    A literal walk of typeface attributes returns +mn-lt and its five
    siblings, so the families here have to be real names and none of them
    may be a token.
    """
    report = run("report", str(latin_deck), fonts=licensed_fonts)
    families = report["families_used"]
    assert families, "a default deck still asks for the theme fonts"
    assert not any(name.startswith("+") for name in families)
    assert any("theme" in where for where in families.values())


def test_a_deck_with_no_arabic_embeds_only_what_it_uses(
        latin_deck, licensed_fonts):
    report = run("embed", str(latin_deck), fonts=licensed_fonts)
    assert FAMILY not in report["families_used"]
    assert all(record["family"] != FAMILY for record in report["embedded"])
    for record in report["embedded"]:
        assert record["family"] in report["families_used"]
    assert_consistent(latin_deck)


# ---------------------------------------------------------------------------
# The licence gate
# ---------------------------------------------------------------------------


def test_licence_gate_refuses_a_face_with_no_licence_file(
        deck, unlicensed_fonts):
    report = run("embed", str(deck), fonts=unlicensed_fonts)
    assert report["embedded"] == []
    refused = [s for s in report["skipped"] if s["family"] == FAMILY]
    assert refused, "the unlicensed family should have been refused"
    assert refused[0]["reason"] == "no licence file found"
    assert any("--allow-unlicensed" in line for line in refused[0]["detail"])
    assert "ppt/fonts/font1.fntdata" not in package(deck)["names"]


def test_the_override_flag_embeds_the_same_refused_face(
        deck, unlicensed_fonts):
    report = run("embed", str(deck), "--allow-unlicensed",
                 fonts=unlicensed_fonts)
    record = [r for r in report["embedded"] if r["family"] == FAMILY]
    assert record, "the override should have let the face through"
    assert record[0]["licence"] is None
    assert record[0]["licence_override"] is True
    assert_consistent(deck)


def test_a_licence_beside_the_font_passes_the_gate_and_is_reported(
        deck, licensed_fonts):
    report = run("embed", str(deck), fonts=licensed_fonts)
    record = [r for r in report["embedded"] if r["family"] == FAMILY][0]
    assert record["licence_override"] is False
    assert record["licence"].endswith("OFL.txt")
    assert Path(record["licence"]).is_file()


def test_a_licence_above_the_font_still_counts(tmp_path, ttf, deck):
    """Distributions file one licence for a whole family tree.

    Walking up as far as the font root finds it. Walking further would
    let an unrelated LICENSE in the home directory vouch for a font, so
    the walk stops there.
    """
    root = tmp_path / "tree"
    (root / FAMILY / "static").mkdir(parents=True)
    shutil.copy(ttf, root / FAMILY / "static" / f"{FAMILY}-Regular.ttf")
    (root / FAMILY / "OFL.txt").write_text("OFL 1.1\n", encoding="utf-8")
    report = run("embed", str(deck), fonts=root)
    record = [r for r in report["embedded"] if r["family"] == FAMILY][0]
    assert record["licence"] == str(root / FAMILY / "OFL.txt")


# ---------------------------------------------------------------------------
# Missing faces
# ---------------------------------------------------------------------------


def test_a_missing_face_is_reported_and_not_fatal(deck, licensed_fonts):
    report = run("embed", str(deck), "--family", "No Such Family At All",
                 "--allow-unlicensed", fonts=licensed_fonts)
    assert report["ok"] is True
    missing = [s for s in report["skipped"]
               if s["family"] == "No Such Family At All"]
    assert missing and missing[0]["reason"] == "no font file found"
    assert reopens(deck)


def test_one_missing_face_does_not_cost_the_others(deck, licensed_fonts):
    report = run("embed", str(deck), "--family", "No Such Family At All",
                 "--family", FAMILY, fonts=licensed_fonts)
    assert [r["family"] for r in report["embedded"]] == [FAMILY]
    assert [s["family"] for s in report["skipped"]] == ["No Such Family At All"]
    assert_consistent(deck)


# ---------------------------------------------------------------------------
# The package after the pass
# ---------------------------------------------------------------------------


def test_the_package_stays_valid_and_the_deck_still_opens(
        deck, licensed_fonts):
    titles_before = reopens(deck)
    report = run("embed", str(deck), fonts=licensed_fonts)
    assert report["bytes_added"] > 0
    assert report["size_after"] > report["size_before"]
    assert_consistent(deck)
    assert reopens(deck) == titles_before
    assert run("verify", str(deck))["problems"] == []


def test_regular_and_bold_land_in_separate_parts(deck, licensed_fonts):
    """PowerPoint tells the two apart by the part, not by the bytes."""
    report = run("embed", str(deck), fonts=licensed_fonts)
    slots = [r for r in report["embedded"] if r["family"] == FAMILY][0]["slots"]
    assert set(slots) == {"regular", "bold"}
    assert slots["regular"]["part"] != slots["bold"]["part"]
    entries = embedded_entries(package(deck)["presentation"])
    assert entries[FAMILY]["regular"] != entries[FAMILY]["bold"]


def test_the_embed_flags_are_set_on_the_presentation(deck, licensed_fonts):
    run("embed", str(deck), fonts=licensed_fonts)
    presentation = package(deck)["presentation"]
    assert presentation.get("embedTrueTypeFonts") == "1"
    # Full faces are embedded, not subsets, and claiming otherwise
    # confuses the editability prompt PowerPoint shows the reader.
    assert presentation.get("saveSubsetFonts") == "0"


def test_embedding_is_idempotent(deck, licensed_fonts):
    run("embed", str(deck), fonts=licensed_fonts)
    first = package(deck)
    first_entries = embedded_entries(first["presentation"])

    second_report = run("embed", str(deck), fonts=licensed_fonts)
    assert second_report["bytes_added"] == 0
    second = package(deck)

    assert second["names"] == first["names"]
    assert second["overrides"] == first["overrides"]
    assert second["targets"] == first["targets"]
    assert embedded_entries(second["presentation"]) == first_entries
    assert_consistent(deck)


def test_a_deck_with_nothing_to_embed_grows_no_empty_element(
        deck, unlicensed_fonts):
    """An empty p:embeddedFontLst is not valid, so it must not be left."""
    run("embed", str(deck), fonts=unlicensed_fonts)
    presentation = package(deck)["presentation"]
    assert presentation.find(f"{{{P_NS}}}embeddedFontLst") is None
    assert reopens(deck)


def test_output_flag_leaves_the_source_untouched(deck, licensed_fonts,
                                                 tmp_path):
    out = tmp_path / "shipped.pptx"
    run("embed", str(deck), "-o", str(out), fonts=licensed_fonts)
    assert "ppt/fonts/font1.fntdata" in package(out)["names"]
    assert "ppt/fonts/font1.fntdata" not in package(deck)["names"]


# ---------------------------------------------------------------------------
# Real faces on this machine
# ---------------------------------------------------------------------------


def _family_is_findable(family: str) -> bool:
    sys.path.insert(0, str(CLI.parent))
    try:
        import pptx_embed_fonts
    except ImportError:
        return False
    finally:
        sys.path.pop(0)
    faces, _ = pptx_embed_fonts.find_faces(family)
    return bool(faces)


@pytest.mark.skipif(not _family_is_findable(ARABIC_FAMILY),
                    reason=f"{ARABIC_FAMILY} is not installed as .ttf")
def test_an_arabic_deck_embeds_its_arabic_face(tmp_path):
    """The case the feature exists for.

    Without the embedded face this deck renders as whatever the reader's
    machine substitutes for the family name, or as tofu.
    """
    arabic = _deck(tmp_path, "arabic.pptx", ARABIC_FAMILY, ARABIC_TEXT)
    report = run("embed", str(arabic))
    assert ARABIC_FAMILY in report["families_used"]
    record = [r for r in report["embedded"]
              if r["family"] == ARABIC_FAMILY]
    assert record, report["skipped"]
    assert "regular" in record[0]["slots"]
    assert_consistent(arabic)
    assert reopens(arabic) == [ARABIC_TEXT]


PLEX_ARABIC_LICENCE = Path(
    "/usr/share/licenses/ibm-plex-sans-arabic-fonts/license.txt")


@pytest.mark.skipif(not PLEX_ARABIC_LICENCE.is_file()
                    or not _family_is_findable(ARABIC_FAMILY),
                    reason="the packaged IBM Plex Arabic licence is not here")
def test_the_licence_is_found_through_another_file_of_the_family(tmp_path):
    """A licence covers the font software, not one copy of one file.

    On this machine the .ttf copies of IBM Plex Sans Arabic sit in an
    unpackaged directory under the home while the .otf copies carry the
    rpm licence. Searching from the embedded file alone would refuse the
    exact face Arabic decks are set in.
    """
    arabic = _deck(tmp_path, "arabic.pptx", ARABIC_FAMILY, ARABIC_TEXT)
    report = run("report", str(arabic))
    record = [r for r in report["embedded"]
              if r["family"] == ARABIC_FAMILY]
    assert record, report["skipped"]
    assert record[0]["licence_override"] is False
    assert Path(record[0]["licence"]).is_file()


def test_a_longer_family_name_is_not_mistaken_for_this_one(
        tmp_path, ttf, deck):
    """IBM Plex Sans must not claim IBMPlexSansArabic-Regular.ttf."""
    root = tmp_path / "lookalike"
    (root / "other").mkdir(parents=True)
    shutil.copy(ttf, root / "other" / f"{FAMILY}Extended-Regular.ttf")
    (root / "other" / "OFL.txt").write_text("OFL 1.1\n", encoding="utf-8")
    report = run("embed", str(deck), fonts=root)
    assert all(r["family"] != FAMILY for r in report["embedded"])
    refused = [s for s in report["skipped"] if s["family"] == FAMILY]
    assert refused and refused[0]["reason"] == "no font file found"


@pytest.mark.skipif(not _family_is_findable("Cairo"),
                    reason="Cairo is not installed as .ttf")
def test_a_variable_font_fills_one_slot_and_says_so(deck):
    """One file covers the whole weight axis.

    PowerPoint renders an embedded face at its default instance, so
    pointing regular and bold at the same bytes would ship the same face
    twice under two labels for no gain.
    """
    report = run("embed", str(deck), "--family", "Cairo",
                 "--allow-unlicensed")
    record = [r for r in report["embedded"] if r["family"] == "Cairo"]
    if not record:
        pytest.skip("Cairo is installed as static faces on this machine")
    slots = record[0]["slots"]
    paths = [slot["file"] for slot in slots.values()]
    assert len(paths) == len(set(paths)), "a file was used for two slots"
    if "variable" in " ".join(record[0]["notes"]).lower():
        assert set(slots) == {"regular"}
    assert_consistent(deck)
