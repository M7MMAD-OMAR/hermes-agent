"""Choosing a layout from content, and splitting what does not fit.

Two failures this guards against, and they are the ones a reader sees
first. A deck where every slide is a title over a bullet list, because
the spec author reached for the only layout they remembered. And a slide
that clips, because the builder returned exactly one slide per section
no matter how much content the section held.

The rule the tests pin: a layout that cannot hold its content is never
built broken. It is split across slides, or degraded to a layout that
reads, and either way the change is reported rather than silent.
"""

from __future__ import annotations

import importlib.util
import json
import subprocess
import sys
from pathlib import Path

import pytest

REPO = Path(__file__).resolve().parents[2]
SCRIPTS = REPO / "skills" / "productivity" / "powerpoint" / "scripts"
HOUSE = REPO / "skills" / "productivity" / "house-style" / "scripts"
CREATE = SCRIPTS / "pptx_create.py"


@pytest.fixture(scope="module")
def create():
    for path in (SCRIPTS, HOUSE):
        if str(path) not in sys.path:
            sys.path.insert(0, str(path))
    spec = importlib.util.spec_from_file_location("pptx_create_under_test",
                                                  CREATE)
    module = importlib.util.module_from_spec(spec)
    sys.modules["pptx_create_under_test"] = module
    spec.loader.exec_module(module)
    return module


@pytest.fixture(scope="module")
def house():
    if str(HOUSE) not in sys.path:
        sys.path.insert(0, str(HOUSE))
    import house_style  # noqa: PLC0415

    return house_style


@pytest.fixture
def style(create, house):
    pytest.importorskip("pptx")
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    return create.DeckStyle(prs, house.load_theme(), house)


def plan(create, style, house, spec):
    pages, notes = create.plan_slides(spec, style, house)
    return pages, notes


def build(tmp_path, slides, name="deck") -> dict:
    spec = tmp_path / f"{name}.json"
    spec.write_text(json.dumps({"slides": slides}, ensure_ascii=False),
                    encoding="utf-8")
    out = tmp_path / f"{name}.pptx"
    run = subprocess.run([sys.executable, str(CREATE), str(spec), str(out)],
                         capture_output=True, text=True)
    assert run.returncode == 0, run.stderr
    report = json.loads(run.stdout)
    report["path"] = out
    return report


# ---------------------------------------------------------------------------
# Choosing
# ---------------------------------------------------------------------------


@pytest.mark.parametrize("spec,expected", [
    ({"stats": [{"value": "34%", "label": "growth"}]}, "stat"),
    ({"steps": [{"label": "May"}, {"label": "June"}, {"label": "July"}]},
     "timeline"),
    ({"left": {"heading": "Before"}, "right": {"heading": "After"}},
     "comparison"),
    ({"cards": [{"label": "One"}, {"label": "Two"}]}, "cards"),
    ({"text": "We took the margin and lost the volume.",
      "attribution": "Head of Sales"}, "quote"),
    ({"text": "One customer now carries a fifth of the book."}, "statement"),
    ({"tables": [{"rows": [["a", "b"]]}]}, "title_content"),
])
def test_auto_reads_the_content(create, spec, expected):
    assert create.choose_layout(spec) == expected


def test_a_few_short_parallel_bullets_become_cards(create):
    """A deck of bullet lists is the tell. Short parallel items are a grid."""
    assert create.choose_layout(
        {"bullets": ["Handles held", "Gaskets fell", "Seals flat"]}) == "cards"
    long_items = ["The notice period moves from thirty days back to sixty "
                  "days for either party and that was the whole argument"]
    assert create.choose_layout({"bullets": long_items * 3}) == "title_content"


# ---------------------------------------------------------------------------
# Validating, and degrading rather than breaking
# ---------------------------------------------------------------------------


def test_a_statement_that_is_a_paragraph_degrades(create):
    long_text = " ".join(["word"] * 29)
    layout, why = create.fit_kind("statement", {"text": long_text})
    assert layout == "title_content"
    assert "one sentence" in why


def test_three_options_are_cards_not_a_comparison(create, style, house):
    spec = {"layout": "comparison", "title": "Three ways", "columns": [
        {"heading": "Hold", "body": "Keep the price."},
        {"heading": "Cut", "body": "Restore volume."},
        {"heading": "Split", "body": "Tier the price."}]}
    pages, notes = plan(create, style, house, spec)
    assert len(pages) == 1, "splitting would strand one option on its own"
    assert pages[0]["layout"] == "cards"
    assert [c["label"] for c in pages[0]["cards"]] == ["Hold", "Cut", "Split"]
    assert [c["body"] for c in pages[0]["cards"]][0] == "Keep the price."
    assert any(n["to"] == "cards" for n in notes)


def test_a_two_step_timeline_becomes_cards_and_keeps_its_text(create, style,
                                                              house):
    spec = {"layout": "timeline", "steps": [{"label": "May", "body": "Prices"},
                                            {"label": "July", "body": "Move"}]}
    pages, notes = plan(create, style, house, spec)
    assert pages[0]["layout"] == "cards"
    assert pages[0]["cards"][0]["label"] == "May"
    assert pages[0]["cards"][0]["body"] == "Prices"
    assert notes and "steps" in notes[0]["why"]


def test_a_layout_with_no_items_falls_back_rather_than_drawing_nothing(create):
    assert create.fit_kind("cards", {"cards": []})[0] == "title_content"
    assert create.fit_kind("stat", {})[0] == "title_content"


def test_content_that_fits_is_left_exactly_as_it_was(create, style, house):
    spec = {"layout": "cards", "title": "Three", "cards": [
        {"label": "A"}, {"label": "B"}, {"label": "C"}]}
    pages, notes = plan(create, style, house, spec)
    assert len(pages) == 1
    assert notes == [], "a slide that fits is not worth a note"
    assert pages[0]["cards"] == spec["cards"]


# ---------------------------------------------------------------------------
# Splitting
# ---------------------------------------------------------------------------


def test_six_cards_become_two_slides_of_three(create, style, house):
    cards = [{"label": f"Card {n}", "body": "One line."} for n in range(6)]
    pages, notes = plan(create, style, house,
                        {"layout": "cards", "title": "Where it went",
                         "cards": cards})
    assert [len(p["cards"]) for p in pages] == [3, 3]
    assert pages[1]["title"].endswith("(continued)")
    assert any("split" in n["why"] for n in notes)


def test_a_lonely_trailing_item_is_pulled_back(create):
    """Five cards over pages of four would leave one on its own."""
    pages = create._chunk_balanced([1, 2, 3, 4, 5], 4)
    assert [len(p) for p in pages] == [3, 2]


def test_a_long_bullet_list_splits_on_the_reading_budget(create, style, house,
                                                         tmp_path):
    bullets = ["The notice period moves from thirty days back to sixty days "
               "for either party" for _ in range(8)]
    pages, _ = plan(create, style, house,
                    {"layout": "title_content", "title": "Clauses",
                     "bullets": bullets})
    limits = house.LIMITS
    assert len(pages) > 1
    for page in pages:
        words = sum(len(b.split()) for b in page["bullets"])
        assert len(page["bullets"]) <= limits["slide_bullets"]
        assert words <= limits["slide_words_max"], \
            "geometry is not the only budget, reading is the other one"


def test_a_chart_belongs_to_the_first_page_only(create, style, house):
    bullets = ["A clause that runs on for a dozen words or more, easily"
               for _ in range(8)]
    pages, _ = plan(create, style, house,
                    {"layout": "title_content", "title": "Terms",
                     "bullets": bullets,
                     "charts": [{"type": "bar", "categories": ["Q1"],
                                 "series": {"Revenue": [1]}}]})
    assert len(pages) > 1
    assert "charts" in pages[0]
    assert all("charts" not in page for page in pages[1:]), \
        "a chart repeated behind every continuation is a chart nobody reads"


def test_an_arabic_continuation_says_so_in_arabic(create, style, house):
    cards = [{"label": f"بطاقة {n}", "body": "سطر واحد."} for n in range(6)]
    pages, _ = plan(create, style, house,
                    {"layout": "cards", "title": "أين ذهب الربع",
                     "cards": cards})
    assert len(pages) == 2
    assert pages[1]["title"].endswith("(تتمة)")


# ---------------------------------------------------------------------------
# End to end, through the script
# ---------------------------------------------------------------------------


def test_the_report_names_every_change_it_made(tmp_path):
    pytest.importorskip("pptx")
    report = build(tmp_path, [
        {"layout": "auto", "title": "Numbers",
         "stats": [{"value": f"{n}%", "label": f"Metric {n}"}
                   for n in range(5)]},
        {"layout": "statement", "title": "Risk",
         "text": " ".join(["word"] * 30)},
    ])
    assert report["slides"] == 3, "five stats do not fit one slide"
    plan_notes = report["plan"]
    assert {n["to"] for n in plan_notes} >= {"stat", "title_content"}
    assert all("section" in n for n in plan_notes)


def test_a_deck_that_fits_reports_no_plan_at_all(tmp_path):
    pytest.importorskip("pptx")
    report = build(tmp_path, [
        {"layout": "title", "title": "Q3 review", "subtitle": "Sales"},
        {"layout": "title_content", "title": "Two points",
         "bullets": ["Handles held", "Gaskets fell"]},
    ], name="quiet")
    assert report["slides"] == 2
    assert "plan" not in report, "nothing was changed, so nothing is reported"


def test_the_split_deck_still_passes_the_geometry_lint(tmp_path):
    pytest.importorskip("pptx")
    report = build(tmp_path, [
        {"layout": "cards", "title": "Six",
         "cards": [{"label": f"Card {n}", "body": "One short line."}
                   for n in range(6)]},
        {"layout": "timeline", "title": "Seven steps",
         "steps": [{"label": f"M{n}", "body": "Something happened"}
                   for n in range(7)]},
    ], name="split")
    lint = subprocess.run(
        [sys.executable, str(HOUSE / "style_lint.py"), str(report["path"]),
         "--only", "geometry", "--json"], capture_output=True, text=True)
    findings = json.loads(lint.stdout)["findings"]
    assert findings == [], f"splitting produced a broken slide: {findings}"


# ---------------------------------------------------------------------------
# Growing, which is fitting in the other direction
# ---------------------------------------------------------------------------


def test_a_short_statement_grows_into_the_room_it_was_given(create, style):
    short = "One customer carries a fifth of the book."
    long_one = ("One customer now carries a fifth of the book and the "
                "contract renews in March, so the negotiation starts in "
                "January and the whole quarter turns on it.")
    base = style.type["section"]

    grown = create._grow_to_fill(style, short, base, style.content_w, 6.0,
                                 1.15, style.type["cover_title"])
    assert grown > base, "a one line statement in a full slide reads as lost"
    assert grown in set(style.type.values()), \
        "growing means the next role up the scale, not an invented size"
    assert grown <= style.type["cover_title"]

    assert create._grow_to_fill(style, long_one, base, style.content_w, 6.0,
                                1.15, style.type["cover_title"]) == base, \
        "text that already fills the frame is left alone"


def test_growing_is_capped_and_never_leaves_the_scale(create, style, tmp_path):
    pytest.importorskip("pptx")
    from pptx import Presentation

    report = build(tmp_path, [
        {"layout": "statement", "kicker": "THE POINT", "text": "Two words."},
        {"layout": "quote", "text": "We took the margin.",
         "attribution": "Head of Sales"},
    ], name="grown")
    sizes = {round(run.font.size.pt)
             for slide in Presentation(str(report["path"])).slides
             for shape in slide.shapes if shape.has_text_frame
             for para in shape.text_frame.paragraphs for run in para.runs
             if run.font.size}
    roles = set(style.type.values())
    assert sizes <= roles, f"off scale sizes reached the deck: {sizes - roles}"
