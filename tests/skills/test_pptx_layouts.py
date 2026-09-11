"""Contracts for the composed slide layouts in the PowerPoint skill.

The deck path used to offer six placeholder layouts, so every generated
deck was a title over a bullet list repeated. These seven compose shapes
on the house grid instead, which means nothing catches a regression
except geometry: every box has to land inside the margins, no two boxes
that carry text may overlap, and an Arabic slide has to mirror its
columns rather than merely right-align them.

The other half of the contract is that the skill still works installed
alone, so each layout is also built with no theme at all.
"""

from __future__ import annotations

import importlib.util
import sys
from pathlib import Path

import pytest

pytest.importorskip("pptx")

from pptx import Presentation  # noqa: E402
from pptx.util import Inches  # noqa: E402

REPO = Path(__file__).resolve().parents[2]
SCRIPTS = REPO / "skills" / "productivity" / "powerpoint" / "scripts"
HOUSE = REPO / "skills" / "productivity" / "house-style" / "scripts"

EMU = 914400
EPS = 0.02  # inches, slack for float division on the 4 pt grid


def _load(name: str, path: Path):
    spec = importlib.util.spec_from_file_location(name, path)
    module = importlib.util.module_from_spec(spec)
    sys.modules[name] = module
    spec.loader.exec_module(module)
    return module


@pytest.fixture(scope="module")
def pc():
    # The script imports its siblings by bare name, the way it is run.
    for path in (SCRIPTS, HOUSE):
        if str(path) not in sys.path:
            sys.path.insert(0, str(path))
    return _load("pptx_create_under_test", SCRIPTS / "pptx_create.py")


def _deck(widescreen=True):
    prs = Presentation()
    if widescreen:
        prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    else:
        prs.slide_width, prs.slide_height = Inches(10), Inches(7.5)
    return prs


def _build(pc, spec, themed=True, widescreen=True):
    """One slide, built the way main() builds it."""
    prs = _deck(widescreen)
    theme = None
    if themed:
        house = pc.load_house_style()
        theme = house.theme_from_spec({}) if house is not None else None
    style = pc.DeckStyle(prs, theme)
    return prs, pc.build_slide(prs, spec, style), style


def _text_boxes(slide):
    """Every shape carrying visible text, as (left, top, width, height)."""
    boxes = []
    for shape in slide.shapes:
        if not getattr(shape, "has_text_frame", False):
            continue
        if not shape.text_frame.text.strip():
            continue
        boxes.append((shape.left / EMU, shape.top / EMU,
                      shape.width / EMU, shape.height / EMU))
    return boxes


def _overlap(a, b):
    wide = min(a[0] + a[2], b[0] + b[2]) - max(a[0], b[0])
    tall = min(a[1] + a[3], b[1] + b[3]) - max(a[1], b[1])
    return max(0.0, wide) * max(0.0, tall)


def _assert_laid_out(boxes, style):
    assert boxes, "the layout drew nothing"
    right = style.width - style.margin_x
    bottom = style.height - style.margin_bottom
    for left, top, width, height in boxes:
        assert left >= style.margin_x - EPS
        assert left + width <= right + EPS
        assert top >= style.margin_top - EPS
        assert top + height <= bottom + EPS
    for i, first in enumerate(boxes):
        for second in boxes[i + 1:]:
            assert _overlap(first, second) <= EPS, (first, second)


def _left_of(slide, needle):
    """Where the shape carrying this text starts.

    An exact match wins over a substring one, because a heading such as
    "قبل" also appears inside the slide title "التهيئة قبل وبعد" and the
    title starts at the margin on every column.
    """
    hits = [shape for shape in slide.shapes
            if getattr(shape, "has_text_frame", False)
            and needle in shape.text_frame.text]
    exact = [shape for shape in hits if shape.text_frame.text.strip() == needle]
    if exact:
        return exact[0].left / EMU
    if hits:
        return hits[0].left / EMU
    raise AssertionError(f"no shape carries {needle!r}")


# ---------------------------------------------------------------------------
# Specs. Each one sits at the top of its layout's budget in house_style
# LIMITS, because a layout that only survives three words is not a layout.
# ---------------------------------------------------------------------------

SPECS = {
    "statement": {
        "layout": "statement", "kicker": "Where we are",
        "text": "Two products carried the quarter and the rest stood still.",
    },
    "stat": {
        "layout": "stat", "title": "The quarter in three numbers",
        "stats": [
            {"value": "34%", "label": "Revenue growth against a flat market",
             "source": "Finance, Q3 close"},
            {"value": "1.2 million", "label": "Active accounts",
             "source": "Product analytics"},
            {"value": "11 days", "label": "Median time to first value",
             "source": "Onboarding funnel"},
        ],
    },
    "cards": {
        "layout": "cards", "title": "Four bets for the fourth quarter",
        "cards": [
            {"label": "Migration", "body": "Move the last legacy accounts."},
            {"label": "Billing", "body": "Retire the second invoicing path."},
            {"label": "Mobile", "body": "Ship the account switcher."},
            {"label": "Docs", "body": "Rewrite the setup guide."},
        ],
    },
    "comparison": {
        "layout": "comparison", "title": "Onboarding, before and after",
        "columns": [
            {"heading": "Before",
             "body": "Nine screens, two emails and a support call before "
                     "anyone saw the product working at all."},
            {"heading": "After",
             "body": "Two screens and a sample project, with the call kept "
                     "for migrations only."},
        ],
    },
    "timeline": {
        "layout": "timeline", "title": "How the quarter lands",
        "steps": [
            {"label": "October", "text": "Pricing ships to new accounts."},
            {"label": "November", "text": "Onboarding reaches everyone."},
            {"label": "December", "text": "Legacy billing retires."},
            {"label": "January", "text": "Migration closes for good."},
        ],
    },
    "quote": {
        "layout": "quote",
        "text": "We stopped asking for the form, and half the tickets "
                "went away.",
        "attribution": "Head of Support",
    },
    "closing": {
        "layout": "closing", "text": "Ship the pricing change in October.",
        "contact": "finance@example.com",
    },
}

ARABIC = {
    "cards": {
        "layout": "cards", "title": "أربعة رهانات للربع الرابع",
        "cards": [
            {"label": "الترحيل", "body": "نقل آخر الحسابات القديمة."},
            {"label": "الفوترة", "body": "إيقاف مسار الفوترة الثاني."},
            {"label": "الجوال", "body": "إطلاق مبدل الحسابات."},
            {"label": "التوثيق", "body": "إعادة كتابة دليل الإعداد."},
        ],
    },
    "comparison": {
        "layout": "comparison", "title": "التهيئة قبل وبعد",
        "columns": [
            {"heading": "قبل", "body": "تسع شاشات ومكالمة دعم."},
            {"heading": "بعد", "body": "شاشتان ومشروع نموذجي."},
        ],
    },
    "timeline": {
        "layout": "timeline", "title": "كيف يسير الربع الرابع",
        "steps": [
            {"label": "أكتوبر", "text": "التسعير الجديد لكل الحسابات."},
            {"label": "نوفمبر", "text": "التهيئة الجديدة تصل الجميع."},
            {"label": "ديسمبر", "text": "إيقاف مسار الفوترة القديم."},
        ],
    },
}


# ---------------------------------------------------------------------------
# One geometry test per layout
# ---------------------------------------------------------------------------


@pytest.mark.parametrize("widescreen", [True, False], ids=["16:9", "4:3"])
@pytest.mark.parametrize("name", sorted(SPECS))
def test_layout_lands_inside_the_margins_without_overlap(pc, name, widescreen):
    """The grid is derived from the live canvas, so 4:3 has to hold too."""
    _, slide, style = _build(pc, SPECS[name], widescreen=widescreen)
    _assert_laid_out(_text_boxes(slide), style)


@pytest.mark.parametrize("name", sorted(SPECS))
def test_layout_composes_without_the_house_style(pc, name):
    """house_common returns None when the design system is not installed."""
    _, slide, style = _build(pc, SPECS[name], themed=False)
    _assert_laid_out(_text_boxes(slide), style)


def test_a_single_stat_and_a_row_of_stats_both_compose(pc):
    _, slide, style = _build(pc, {"layout": "stat", "value": "34%",
                                  "label": "Revenue growth",
                                  "source": "Finance, Q3 close"})
    _assert_laid_out(_text_boxes(slide), style)


def test_a_wrapping_stat_value_gets_room_for_both_lines(pc):
    """The box is sized from an estimate, so a long number must not spill.

    Four stats is the tight case: the number drops to its 60 pt floor in
    the narrowest column the layout allows.
    """
    _, slide, _ = _build(pc, {"layout": "stat", "stats": [
        {"value": "1.2 million", "label": "Active accounts"},
        {"value": "34%", "label": "Revenue growth"},
        {"value": "11 days", "label": "Median time to first value"},
        {"value": "9.4x", "label": "Return on the migration"}]})
    values = [box for box in _text_boxes(slide) if box[3] > 0.9]
    assert values, "no value box found"
    assert len({round(box[3], 3) for box in values}) == 1, \
        "the columns must share one value height"
    assert min(box[3] for box in values) > 60 * 1.05 / 72 * 1.5, \
        "a value that wraps was given one line of room"


@pytest.mark.parametrize("name", ["cards", "comparison", "timeline"])
def test_a_layout_missing_its_content_still_draws_its_title(pc, name):
    _, slide, _ = _build(pc, {"layout": name, "title": "What changed"})
    text = " ".join(shape.text_frame.text for shape in slide.shapes
                    if getattr(shape, "has_text_frame", False))
    assert "What changed" in text


def test_two_cards_sit_in_one_row_and_four_sit_two_by_two(pc):
    _, two, _ = _build(pc, {"layout": "cards", "cards": [
        {"label": "One", "body": "First."}, {"label": "Two", "body": "Second."}]})
    tops = {round(box[1], 3) for box in _text_boxes(two)}
    assert len(tops) == 1, "two cards must share one row"

    _, four, _ = _build(pc, SPECS["cards"])
    card_tops = {round(box[1], 3) for box in _text_boxes(four)
                 if box[1] > 1.5}
    assert len(card_tops) == 2, "four cards must fall into two rows"


def test_cards_in_a_row_are_the_same_height(pc):
    _, slide, _ = _build(pc, SPECS["cards"])
    heights = {round(box[3], 3) for box in _text_boxes(slide) if box[1] > 1.5}
    assert len(heights) == 1, heights


def test_the_timeline_rule_is_a_hairline_not_a_filled_bar(pc):
    _, slide, _ = _build(pc, SPECS["timeline"])
    rules = [s for s in slide.shapes if s.name == "house rule"]
    assert len(rules) == 1
    assert rules[0].height == 0, "a rule with height would sit under the dots"


# ---------------------------------------------------------------------------
# Mirroring
# ---------------------------------------------------------------------------


def test_arabic_cards_reverse_their_order(pc):
    _, latin, _ = _build(pc, SPECS["cards"])
    assert _left_of(latin, "Migration") < _left_of(latin, "Billing")

    _, arabic, _ = _build(pc, ARABIC["cards"])
    assert _left_of(arabic, "الترحيل") > _left_of(arabic, "الفوترة")
    # The second row mirrors too, not only the first.
    assert _left_of(arabic, "الجوال") > _left_of(arabic, "التوثيق")


def test_arabic_comparison_columns_reverse_their_order(pc):
    _, latin, _ = _build(pc, SPECS["comparison"])
    assert _left_of(latin, "Before") < _left_of(latin, "After")

    _, arabic, style = _build(pc, ARABIC["comparison"])
    assert _left_of(arabic, "قبل") > _left_of(arabic, "بعد")
    _assert_laid_out(_text_boxes(arabic), style)


def test_arabic_timeline_steps_reverse_their_order(pc):
    _, latin, _ = _build(pc, SPECS["timeline"])
    assert _left_of(latin, "October") < _left_of(latin, "January")

    _, arabic, style = _build(pc, ARABIC["timeline"])
    assert _left_of(arabic, "أكتوبر") > _left_of(arabic, "ديسمبر")
    _assert_laid_out(_text_boxes(arabic), style)


def test_arabic_text_aligns_to_the_right(pc):
    from pptx.enum.text import PP_ALIGN

    _, slide, _ = _build(pc, {"layout": "statement",
                              "text": "منتجان حملا الربع بأكمله."})
    alignments = {para.alignment
                  for shape in slide.shapes
                  if getattr(shape, "has_text_frame", False)
                  for para in shape.text_frame.paragraphs}
    assert alignments == {PP_ALIGN.RIGHT}


# ---------------------------------------------------------------------------
# The six stock layouts are untouched
# ---------------------------------------------------------------------------


def test_an_unknown_layout_falls_back_to_title_and_body(pc):
    """The old behaviour, not an exception: a typo must still make a slide."""
    _, slide, _ = _build(pc, {"layout": "spaceship", "title": "Agenda",
                              "bullets": ["Top item", "Second item"]})
    assert slide.shapes.title is not None
    assert slide.shapes.title.text == "Agenda"
    body = " ".join(shape.text_frame.text for shape in slide.shapes
                    if getattr(shape, "has_text_frame", False))
    assert "Top item" in body and "Second item" in body


@pytest.mark.parametrize("name", ["title", "title_content", "section",
                                  "two_content", "title_only", "blank"])
def test_the_stock_layouts_still_use_their_placeholders(pc, name):
    _, slide, _ = _build(pc, {"layout": name, "title": "Agenda",
                              "bullets": ["Top item"]})
    assert slide.slide_layout == slide.part.package.presentation_part \
        .presentation.slide_layouts[pc.LAYOUTS[name]]


def test_a_composed_layout_keeps_notes_and_background(pc):
    _, slide, _ = _build(pc, {"layout": "quote", "text": "Short.",
                              "notes": "Say it slowly.",
                              "background": "1F2937"})
    assert slide.notes_slide.notes_text_frame.text == "Say it slowly."
    assert str(slide.background.fill.fore_color.rgb) == "1F2937"
