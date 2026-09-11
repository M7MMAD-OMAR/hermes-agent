"""Contracts for the cross format document inspector.

The inspector is what an agent reads before it edits somebody else's
file, so the tests that matter are the ones that keep it honest: it must
report what is actually in the package, and it must say "I could not
look" rather than "there is nothing there" when a helper is missing.
"""

from __future__ import annotations

import importlib.util
import json
import subprocess
import sys
from pathlib import Path

import pytest

REPO = Path(__file__).resolve().parents[2]
HOUSE = REPO / "skills" / "productivity" / "house-style" / "scripts"


@pytest.fixture(scope="module")
def inspector():
    if str(HOUSE) not in sys.path:
        sys.path.insert(0, str(HOUSE))
    spec = importlib.util.spec_from_file_location(
        "office_inspect_under_test", HOUSE / "office_inspect.py")
    module = importlib.util.module_from_spec(spec)
    sys.modules["office_inspect_under_test"] = module
    spec.loader.exec_module(module)
    return module


def _docx(tmp_path, arabic=False):
    pytest.importorskip("docx")
    from docx import Document

    doc = Document()
    doc.add_heading("الملخص" if arabic else "Summary", 1)
    doc.add_paragraph("نص عربي كامل للتجربة" if arabic else "Body text here.")
    table = doc.add_table(rows=2, cols=3)
    table.cell(0, 0).text = "Product"
    out = tmp_path / ("ar.docx" if arabic else "report.docx")
    doc.save(str(out))
    return out


def test_word_report_carries_structure_and_targets(inspector, tmp_path):
    report = inspector.inspect(_docx(tmp_path), lint=False)
    assert report["kind"] == "word"
    assert report["summary"]["tables"] == 1
    assert report["summary"]["headings"] == 1
    assert any(s["style"] == "Heading 1" for s in report["structure"])
    # Every target is addressable: that is the point of the section.
    assert all("address" in t for t in report["targets"])
    assert any(t["what"] == "table" for t in report["targets"])


def test_word_report_notices_arabic(inspector, tmp_path):
    report = inspector.inspect(_docx(tmp_path, arabic=True), lint=False)
    assert report["summary"]["arabic"] is True


def test_deck_report_reads_slides_and_charts(inspector, tmp_path):
    pytest.importorskip("pptx")
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.util import Inches

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Revenue"
    slide.placeholders[1].text_frame.text = "Handles carried the quarter"
    data = CategoryChartData()
    data.categories = ["Q1", "Q2"]
    data.add_series("Revenue", (10, 20))
    slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1),
                           Inches(2), Inches(4), Inches(3), data)
    out = tmp_path / "deck.pptx"
    prs.save(str(out))

    report = inspector.inspect(out, lint=False)
    assert report["summary"]["slides"] == 1
    assert report["summary"]["charts"] == 1
    chart = report["inventory"]["charts"][0]
    assert chart["series"] == ["Revenue"]
    assert chart["categories"] == ["Q1", "Q2"]
    assert report["structure"][0]["title"] == "Revenue"


def test_workbook_report_counts_formulas_and_reads_the_header(inspector,
                                                              tmp_path):
    pytest.importorskip("openpyxl")
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.append(["Region", "Revenue"])
    ws.append(["North", 100])
    ws["B3"] = "=SUM(B2:B2)"
    out = tmp_path / "book.xlsx"
    wb.save(str(out))

    report = inspector.inspect(out, lint=False)
    assert report["summary"]["sheets"] == 1
    assert report["summary"]["formulas"] == 1
    assert report["structure"][0]["header"] == ["Region", "Revenue"]


def test_media_inventory_reads_the_package_not_the_object_model(inspector,
                                                                tmp_path):
    """A picture is found by its part, so nothing hides in a text box."""
    pytest.importorskip("docx")
    pytest.importorskip("PIL")
    from docx import Document
    from docx.shared import Mm
    from PIL import Image

    png = tmp_path / "dot.png"
    Image.new("RGB", (40, 24), (180, 72, 46)).save(png)
    doc = Document()
    doc.add_picture(str(png), width=Mm(30))
    out = tmp_path / "with-image.docx"
    doc.save(str(out))

    media = inspector.inspect(out, lint=False)["inventory"]["media"]
    assert len(media) == 1
    assert media[0]["kind"] == "png"
    assert media[0]["width"] == 40 and media[0]["height"] == 24


def test_a_missing_helper_says_so_instead_of_reporting_no_comments(inspector,
                                                                   tmp_path,
                                                                   monkeypatch):
    """The failure this section exists to prevent, reproduced three ways.

    Asserting it against a checkout where every helper is installed tests
    the success branch and nothing else, which is how a silent empty list
    would have shipped green.
    """
    doc = _docx(tmp_path)

    # 1. the helper is not installed at all
    monkeypatch.setattr(inspector, "_skill_script", lambda skill, name: None)
    comments = inspector.inspect(doc, lint=False)["review"]["comments"]
    assert comments["items"] == []
    assert "not installed" in comments["source"]
    assert "data" not in comments, "an absent helper must not look like a read"

    # 2. the helper is there and exits non-zero
    angry = tmp_path / "angry.py"
    angry.write_text("import sys\nsys.stderr.write('boom')\nsys.exit(3)\n",
                     encoding="utf-8")
    monkeypatch.setattr(inspector, "_skill_script", lambda s, n: angry)
    comments = inspector.inspect(doc, lint=False)["review"]["comments"]
    assert "exited 3" in comments["source"]
    assert comments["items"] == []
    assert "boom" in comments.get("stderr", "")

    # 3. the helper is there and prints something that is not JSON
    chatty = tmp_path / "chatty.py"
    chatty.write_text("print('all good, no comments here')\n",
                      encoding="utf-8")
    monkeypatch.setattr(inspector, "_skill_script", lambda s, n: chatty)
    comments = inspector.inspect(doc, lint=False)["review"]["comments"]
    assert comments["items"] == []
    assert "no JSON" in comments.get("note", "")


def test_a_present_helper_is_actually_read(inspector, tmp_path, monkeypatch):
    """The other half: a helper that answers must reach the report."""
    speaking = tmp_path / "speaking.py"
    speaking.write_text(
        "import json\nprint(json.dumps({'threads': [{'id': '7'}]}))\n",
        encoding="utf-8")
    monkeypatch.setattr(inspector, "_skill_script", lambda s, n: speaking)
    comments = inspector.inspect(_docx(tmp_path),
                                 lint=False)["review"]["comments"]
    assert comments["data"]["threads"][0]["id"] == "7"


def test_unsupported_extension_is_refused(inspector, tmp_path):
    odd = tmp_path / "notes.md"
    odd.write_text("hello", encoding="utf-8")
    with pytest.raises(ValueError):
        inspector.inspect(odd)


def test_cli_prints_json_and_honours_section(tmp_path):
    out = _docx(tmp_path)
    script = HOUSE / "office_inspect.py"
    run = subprocess.run(
        [sys.executable, str(script), str(out), "--section", "summary",
         "--no-lint"], capture_output=True, text=True)
    assert run.returncode == 0, run.stderr
    payload = json.loads(run.stdout)
    assert set(payload) == {"file", "kind", "summary"}

    missing = subprocess.run(
        [sys.executable, str(script), str(tmp_path / "nope.docx")],
        capture_output=True, text=True)
    assert missing.returncode == 1
